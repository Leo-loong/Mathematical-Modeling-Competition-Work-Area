# -*- coding: utf-8 -*-
"""Mermaid 流程图 + 新增数据图 — 全中文版"""
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
import numpy as np; import pandas as pd; import os; import seaborn as sns
import requests; import base64; import json

DATA = "/Users/kuangping/Desktop/数学建模/Mathematical-Modeling-Competition-Work-Area/03_模拟演练/交接材料/03_结果与图表包/data"
ATTACH = "/Users/kuangping/Desktop/数学建模/Mathematical-Modeling-Competition-Work-Area/03_模拟演练/附件"
OUT = "/Users/kuangping/Desktop/数学建模/Mathematical-Modeling-Competition-Work-Area/03_模拟演练/03_结果与图表/最终图表"
os.makedirs(OUT, exist_ok=True)

def setup():
    for name in ["PingFang SC","Hiragino Sans GB","Microsoft YaHei","SimHei","Arial Unicode MS"]:
        if name in {f.name for f in fm.fontManager.ttflist}:
            plt.rcParams["font.sans-serif"] = [name, "Arial"]
            plt.rcParams["font.family"] = "sans-serif"
            plt.rcParams["axes.unicode_minus"] = False
            return name
    return None
_ = setup()
sns.set_style("whitegrid"); sns.set_context("paper", font_scale=1.0)
plt.rcParams.update({"figure.dpi":150,"savefig.dpi":300,"font.size":10,
    "axes.titlesize":11,"axes.labelsize":9,"legend.fontsize":7.5,
    "xtick.labelsize":8,"ytick.labelsize":8,
    "axes.spines.top":False,"axes.spines.right":False})
C10,C15,CFIT,CGREEN = "#2166AC","#B2182B","#D6604D","#4DAF4A"

def render_mermaid(code, fname, theme="neutral"):
    graph = json.dumps({"code":code,"mermaid":{"theme":theme}})
    encoded = base64.urlsafe_b64encode(graph.encode()).decode()
    url = f"https://mermaid.ink/img/{encoded}?type=png"
    resp = requests.get(url, timeout=30)
    if resp.status_code == 200:
        with open(os.path.join(OUT, fname), "wb") as f: f.write(resp.content)
        print(f"  mermaid: {fname}")
    else: print(f"  FAIL ({resp.status_code}): {fname}")

# ========== 新增数据图 ==========

def draw_extrema():
    """条纹极值检测图"""
    df = pd.read_csv(os.path.join(DATA, "fig3_data.csv"))
    sigma = df["sigma"].values; F = df["F_meas"].values
    from scipy.signal import argrelextrema
    mx = argrelextrema(F, np.greater, order=30)[0]
    mn = argrelextrema(F, np.less, order=30)[0]
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.plot(sigma, F, "-", lw=0.4, color=C10, alpha=0.5, label="条纹分量 F(σ)")
    ax.scatter(sigma[mx], F[mx], c="red", s=25, zorder=5, label="极大值点")
    ax.scatter(sigma[mn], F[mn], c="blue", s=25, zorder=5, label="极小值点")
    ax.text(0.98, 0.95, f"检测到极值点数: {len(mx)+len(mn)}",
            transform=ax.transAxes, fontsize=9, ha="right", va="top",
            bbox=dict(boxstyle="round,pad=0.3", fc="lightyellow", alpha=0.8))
    ax.set_xlabel("波数 σ (cm⁻¹)"); ax.set_ylabel("条纹分量")
    ax.legend(fontsize=8, loc="upper right")
    fig.tight_layout()
    fig.savefig(os.path.join(OUT, "figA_extrema.png"), dpi=300, pad_inches=0.08)
    plt.close(fig); print("  figA done")

def draw_si():
    """Si 反射率谱"""
    import openpyxl
    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(8, 5.5), sharex=True)
    for fname, ax, color, label in [("附件3.xlsx",ax1,"#1B9E77","10°"),("附件4.xlsx",ax2,"#D95F02","15°")]:
        wb = openpyxl.load_workbook(os.path.join(ATTACH, fname)); ws = wb.active
        sig, Rr = [], []
        for r in range(2, ws.max_row+1):
            sv, rv = ws.cell(r,1).value, ws.cell(r,2).value
            if sv and rv: sig.append(float(sv)); Rr.append(float(rv))
        s = np.array(sig); R = np.array(Rr)
        ax.plot(s, R, "-", lw=0.35, color=color, alpha=0.85, label=f"实测谱 ({label})")
        ax.fill_between(s, 0, R, alpha=0.06, color=color)
        ax.set_ylabel("反射率 R (%)"); ax.legend(loc="upper right", fontsize=8)
    ax2.set_xlabel("波数 σ (cm⁻¹)")
    fig.tight_layout()
    fig.savefig(os.path.join(OUT, "figB_si_spectrum.png"), dpi=300, pad_inches=0.08)
    plt.close(fig); print("  figB done")

def draw_dual():
    """双角度条纹对比"""
    df = pd.read_csv(os.path.join(DATA, "fig2_data.csv"))
    mask = (df["sigma10"]>=1800) & (df["sigma10"]<=2100)
    s10 = df.loc[mask,"sigma10"].values; s15 = df.loc[mask,"sigma15"].values
    F10 = df.loc[mask,"R10"].values - df.loc[mask,"bg10"].values
    F15 = df.loc[mask,"R15"].values - df.loc[mask,"bg15"].values
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.plot(s10, F10, "-", lw=1.2, color=C10, alpha=0.9, label="条纹分量 10°")
    ax.plot(s15, F15, "--", lw=1.2, color=C15, alpha=0.9, label="条纹分量 15°")
    ax.set_xlabel("波数 σ (cm⁻¹)"); ax.set_ylabel("条纹分量（扣除背景后）")
    ax.legend(fontsize=8, loc="upper right")
    fig.tight_layout()
    fig.savefig(os.path.join(OUT, "figC_dual_angle.png"), dpi=300, pad_inches=0.08)
    plt.close(fig); print("  figC done")

# ========== Mermaid 流程图（全中文） ==========

MERMAID = {
    "figM1_algorithm_flow": """
flowchart TD
    A["附件1/2<br/>R(σ) 10°、15°"] --> B["预处理<br/>波数轴校验<br/>剩余射线带剔除<br/>poly7 背景扣除<br/>噪声估计 ≈0.012%"]
    B --> C1["E1 极值回归<br/>光谱域<br/>条纹极值→线性回归<br/>d=k/(4n₁cosθ₁)"]
    B --> C2["E2 谱域 FFT<br/>二次频域<br/>Hann+8倍零填充<br/>峰位→厚度"]
    B --> C3["E3 全谱拟合<br/>模型域<br/>v1.0 正演非线性LS<br/>初值由E1/E2供给"]
    C1 --> D["逆方差加权融合<br/>wᵢ=(1/σᵢ²)/Σ(1/σⱼ²)"]
    C2 --> D; C3 --> D
    D --> E["L1 检验<br/>估计器间一致性<br/>χ²≤5.99→通过"]
    E --> F["双角度汇总<br/>L2 检验 经验σ<br/>χ²=0.20→通过"]
    F --> G["最终结果<br/>d=8.03±0.05 μm<br/>系统区间[8.0,8.9]μm"]

    style A fill:#E3F2FD,stroke:#1565C0
    style B fill:#E8F5E9,stroke:#2E7D32
    style C1 fill:#FFF3E0,stroke:#E65100
    style C2 fill:#FFF3E0,stroke:#E65100
    style C3 fill:#FFF3E0,stroke:#E65100
    style D fill:#F3E5F5,stroke:#7B1FA2
    style E fill:#E0F2F1,stroke:#00695C
    style F fill:#E0F2F1,stroke:#00695C
    style G fill:#FFEBEE,stroke:#C62828
""",

    "figM2_model_hierarchy": """
flowchart LR
    A["v1.0 双光束模型<br/>━━━━━━<br/>R(σ)=r₀₁²+ρ₁₂²<br/>+2r₀₁ρ₁₂cos(2πpσ+φ₁₂)<br/><br/>常数 n₁=2.55<br/>常数 φ₁₂≈-0.28π<br/>常数 ρ₁₂≈0.165"] -->|"三域互证融合<br/>+ 两级一致性检验"| B["v1.5 融合模型<br/>━━━━━━<br/>E1×E2×E3 逆方差融合<br/>bootstrap N=200<br/>双角度汇总<br/><br/>d=8.03±0.05 μm<br/>系统区间[8.0,8.9]μm"]
    B -->|"退化验证<br/>q→0时Δ=3.9×10⁻¹⁶"| C["v2.0 Airy多光束<br/>━━━━━━<br/>Airy三介质闭式<br/>+ 固定Drude ε(σ)<br/><br/>SiC:多光束可忽略<br/>Drude相位频变主导<br/>修正+0.61→8.64μm<br/><br/>Si:多光束显著<br/>d=3.43μm[3.28,3.58]"]
    A -.->|退化方向| C

    style A fill:#E3F2FD,stroke:#1565C0
    style B fill:#C8E6C9,stroke:#2E7D32
    style C fill:#FFCDD2,stroke:#C62828
""",

    "figM3_validation_chain": """
flowchart LR
    V1["V1 合成闭环<br/>━━━━<br/>d₀=10,30,100μm<br/>全管线恢复<br/>|Δd|≤0.06μm"]
    V2["V2 双角度互证<br/>━━━━<br/>10° vs 15°<br/>Δd=0.031μm<br/>χ²L₂=0.20 通过"]
    V3["V3 残差诊断<br/>━━━━<br/>std=0.50%<br/>lag-1=0.999<br/>→触发v2.0升级"]
    V4["V4 多算法交叉<br/>━━━━<br/>E1/E2/E3×2角度<br/>max|Δ|≤0.24μm<br/>全部归因无超差"]
    V1 --> V2 --> V3 --> V4
    V4 --> R["可靠性闭环<br/>━━━━<br/>算法无偏→数据自洽<br/>→模型归因→路径互证"]
    style V1 fill:#E8EAF6,stroke:#283593
    style V2 fill:#E8EAF6,stroke:#283593
    style V3 fill:#E8EAF6,stroke:#283593
    style V4 fill:#E8EAF6,stroke:#283593
    style R fill:#C8E6C9,stroke:#2E7D32
""",

    "figM4_effect_separation": """
graph TD
    Q["SiC残差来源判定"] --> M00["M00: 双光束+常数r₁₂<br/>v1.5基准<br/>red-χ²=1862"]
    Q --> M10["M10: 双光束+固定Drude<br/>red-χ²=1432<br/>下降23%"]
    Q --> M01["M01: Airy+常数r₁₂<br/>red-χ²=1836<br/>仅降1.4%"]
    Q --> M11["M11: Airy+固定Drude<br/>red-χ²=1436"]
    M00 --> C1["结论1<br/>多光束可忽略<br/>q≈0.002-0.005"]
    M01 --> C1
    M10 --> C2["结论2<br/>Drude相位频变主导<br/>修正+0.61μm"]
    M11 --> C2
    style Q fill:#FFF9C4,stroke:#F57F17
    style M00 fill:#E3F2FD,stroke:#1565C0
    style M10 fill:#FFCDD2,stroke:#C62828
    style M01 fill:#E3F2FD,stroke:#1565C0
    style M11 fill:#FFCDD2,stroke:#C62828
    style C1 fill:#C8E6C9,stroke:#2E7D32
    style C2 fill:#FFE0B2,stroke:#E65100
""",
}

if __name__ == "__main__":
    print("Mermaid ...")
    for k, v in MERMAID.items(): render_mermaid(v.strip(), f"{k}.png")
    print("Data figs ...")
    draw_extrema(); draw_si(); draw_dual()
    print("all done")