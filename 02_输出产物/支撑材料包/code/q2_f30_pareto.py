# -*- coding: utf-8 -*-
"""
F-30 精度–效率帕累托图 · 数据导出
=================================
同一算例（0–3 h、Δr=0.25 mm、附件1 时变边界＋H6 延拓）下，比较**时间格式／步长方案**
的（相对误差，墙钟）：

  参考解 = BDF2, n_sub = 1/64 s（最精细，仅作误差基准，不入图）
  方案   = ① 后向欧拉 1/32 s（主力）  ② BDF2 1/8 s  ③ BDF2 1/16 s

相对误差定义：max_r |C(r,t_end) − C_ref(r,t_end)| / max_r C_ref

输出：`04_图表包/data/fig_q2_pareto.csv`（列：scheme, n_sub, wall_s, max_abs_dC, rel_err）
      ＋ `30_图表/03_图数据准备/`（双端同写）

注：E9 的「自适应」方案（3 h 墙钟 9.9 s、与固定解偏差 ≤9.3e-6 kg/kg）**口径不同**
（其偏差是相对固定解、非相对参考解），故**不纳入本图**，仅在 F-30 卡中备注。

用法：python q2_f30_pareto.py            # dry-run
      python q2_f30_pareto.py --go       # 实算（约 3 min）
"""
import io
import os
import sys
import time

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

for _v in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS',
           'NUMEXPR_NUM_THREADS', 'VECLIB_MAXIMUM_THREADS'):
    os.environ.setdefault(_v, '1')

import numpy as np                                                      # noqa: E402
from openpyxl import load_workbook                                      # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.abspath(os.path.join(HERE, '..', '..'))
WS = os.path.abspath(os.path.join(PKG, '..'))
ATT1 = os.path.join(WS, '10_赛题', 'A题', '附件', '附件1.xlsx')
OUT1 = os.path.join(PKG, '04_图表包', 'data')
OUT2 = os.path.join(WS, '30_图表', '03_图数据准备')

sys.path.insert(0, HERE)
import q2_core_c as base                                                # noqa: E402
import q2_core_bdf2 as bdf2mod                                          # noqa: E402

GO = '--go' in sys.argv
N, DR, NSTEP = 80, 2.5e-4, 10800        # 0–3 h
TINF, CINF = 50.00, 0.04999


def load_env():
    wb = load_workbook(ATT1, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    d = [r for r in list(ws.iter_rows(values_only=True))[1:] if r[0] is not None]
    t = np.array([float(r[0]) for r in d])
    T = np.array([float(r[1]) for r in d])
    C = np.array([float(r[2]) for r in d])
    wb.close()
    return t, T, C


def run(n_sub, use_bdf2, t, Tv, Cv):
    Tenv = lambda s: float(np.interp(s, t, Tv)) if s <= t[-1] else TINF       # noqa: E731
    Cenv = lambda s: float(np.interp(s, t, Cv)) if s <= t[-1] else CINF       # noqa: E731
    mod = bdf2mod if use_bdf2 else base
    t0 = time.time()
    res, _, _ = mod.run_q2(N=N, dr=DR, dt_out=1.0, nsteps=NSTEP,
                           Tenv_fn=Tenv, Cenv_fn=Cenv, n_sub=n_sub,
                           mode='coupled', bdf2=use_bdf2) if use_bdf2 else \
               mod.run_q2(N=N, dr=DR, dt_out=1.0, nsteps=NSTEP,
                          Tenv_fn=Tenv, Cenv_fn=Cenv, n_sub=n_sub, mode='coupled')
    return res['C_end'], time.time() - t0


def main():
    t, Tv, Cv = load_env()
    rows = []

    print('参考解：BDF2 n_sub=64 …')
    Cref, w_ref = run(64, True, t, Tv, Cv)
    scale = float(np.max(Cref))
    print('  参考墙钟 %.1f s' % w_ref)

    for label, n_sub, use_b in (('BE-1/32（主力）', 32, False),
                                ('BDF2-1/8', 8, True),
                                ('BDF2-1/16', 16, True)):
        C, w = run(n_sub, use_b, t, Tv, Cv)
        d = float(np.max(np.abs(C - Cref)))
        rel = d / scale
        rows.append('%s,%d,%.2f,%.6e,%.6e' % (label, n_sub, w, d, rel))
        print('  %-16s 墙钟 %6.2f s   max|ΔC| %.3e   相对误差 %.3e' % (label, w, d, rel))

    body = 'scheme,n_sub,wall_s,max_abs_dC,rel_err\n' + '\n'.join(rows) + '\n'
    for d in (OUT1, OUT2):
        os.makedirs(d, exist_ok=True)
        with io.open(os.path.join(d, 'fig_q2_pareto.csv'), 'w',
                     encoding='utf-8', newline='\n') as f:
            f.write(body)
    print('已写 fig_q2_pareto.csv（双端）')


if __name__ == '__main__':
    if not GO:
        print('[SCALE] 将实算 1 个参考解（BDF2 1/64）＋ 3 个方案（0–3 h、N=80），约 3 min；加 --go 执行。')
        sys.exit(0)
    main()
