# -*- coding: utf-8 -*-
r"""
Q3 诊断 ②｜界面 $D$ 取法对照（判定候选 m2：强变系数胞内的界面平均失真）
================================================================================
**背景（解析定位）**：代码的界面系数取**两端点 $D$ 的调和平均** $2D_lD_r/(D_l+D_r)$。
该公式只对"**界面处系数间断**"（分段常数）成立；当 $D$ 随 $C$ **平滑但陡峭**变化时，
胞内真实稳态通量给出的是**沿 $C$ 的积分平均**
$$J=\frac{1}{\Delta r}\int_{C_{i+1}}^{C_i}D(C,T)\,\mathrm dC
  \quad\Longrightarrow\quad
  \langle D\rangle=\frac{1}{\Delta C}\int D\,\mathrm dC,$$
而调和平均 $\approx2\min(D_l,D_r)$，在 $D$ 陡变时**严重低估**（Q3 尾部代表单元估算低估约 73 倍）。

**本实验**：在**同一粗网格**（$\Delta r=0.25$ mm，$N=80$）上，仅更换界面 $D$ 的取法：
  · mode 0：调和平均（**现行口径**，应与 $t_{\rm end}=87.4933$ h 一致 ⟹ 兼作新模块自检）；
  · mode 1：算术平均；
  · mode 2：沿 $C$ 的 8 点积分平均（最接近正确）。
其余口径**完全不变**（附录3、1/32 s、IMEX、事件驱动 $C<0.15$、二分定位）。

**判读**：若 mode 1／2 使 $t_{\rm end}$ 大幅下降并接近"表面加密网格"的 57.6 h ⟹ **m2 确认**，
且说明"粗网格 + 正确界面处理"即可恢复精度（修复路线）。

用法：
  python q3_diag_iface.py --mode 0            # 规模声明
  python q3_diag_iface.py --mode 0 --go       # 执行（或 ALLOW_RUN=1）
"""
import sys
import os
import time
import argparse

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass
for _v in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS'):
    os.environ.setdefault(_v, '1')

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), 'Q2', 'code'))

import openpyxl                                                     # noqa: E402
from q2_core_c import (njit, _rho_nb, _cp_nb, _k_nb, _iface_nb,      # noqa: E402
                       _assemble_T_nb, _rhs_T_nb, _assemble_C_nb,
                       _rhs_C_nb, _thomas_nb, DEF2)
import q3_core as q3                                                # noqa: E402

LOG = os.path.join(HERE, 'q3_diag_iface_log.txt')
WS = os.path.abspath(os.path.join(HERE, '..', '..', '..'))
ATT1 = os.path.join(WS, '10_赛题', 'A题', '附件', '附件1.xlsx')
N, DR = 80, 2.5e-4
HH, DT_OUT = 1.0 / 32.0, 60.0
NSUB = 1920
R0, HC, KM = DEF2['R0'], DEF2['h'], DEF2['km']
C_CRIT = 0.15
D_FAC = 1.0
BUF = []


def say(s=''):
    print(s)
    BUF.append(s)


@njit(cache=True, fastmath=False)
def _D_i(Carr, Tarr, i, d_fac):
    c = Carr[i]
    if c < 1e-9:
        c = 1e-9
    return d_fac * 2.4e-3 * np.exp(-0.45 / c) * np.exp(-3850.0 / (Tarr[i] + 273.15))


@njit(cache=True, fastmath=False)
def _faceD(Carr, Tarr, d_fac, mode, nsamp):
    """界面 D 数组（长度 N）。mode 0=调和、1=算术、2=沿 C 的积分平均。"""
    n = Carr.shape[0] - 1
    out = np.zeros(n)
    for i in range(n):
        dl = _D_i(Carr, Tarr, i, d_fac)
        dr = _D_i(Carr, Tarr, i + 1, d_fac)
        if mode == 0:
            out[i] = 2.0 * dl * dr / (dl + dr) if (dl > 0.0 and dr > 0.0) else 0.0
        elif mode == 1:
            out[i] = 0.5 * (dl + dr)
        else:
            a = Carr[i]
            b = Carr[i + 1]
            Tm = 0.5 * (Tarr[i] + Tarr[i + 1])
            s = 0.0
            for k in range(nsamp):
                x = a + (b - a) * (k + 0.5) / nsamp
                if x < 1e-9:
                    x = 1e-9
                s += d_fac * 2.4e-3 * np.exp(-0.45 / x) * np.exp(-3850.0 / (Tm + 273.15))
            out[i] = s / nsamp
    return out


@njit(cache=True, fastmath=False)
def _step_nb(N, dr, h, n_sub, R0, hc, km, T, C, t0,
             ts_env, Tenv_v, Cenv_v, d_fac, mode, nsamp,
             omega, tol, maxit):
    NC = N + 1
    Tcur = T.copy()
    Ccur = C.copy()
    aT = np.zeros(NC); bT = np.zeros(NC); cT = np.zeros(NC)
    aC = np.zeros(NC); bC = np.zeros(NC); cC = np.zeros(NC)
    dT = np.zeros(NC); dC = np.zeros(NC)
    inner = 0
    resmax = 0.0
    for s in range(n_sub):
        ts = t0 + (s + 1) * h
        tinf = q3._interp1_nb(ts, ts_env, Tenv_v)
        cinf = q3._interp1_nb(ts, ts_env, Cenv_v)
        rho = _rho_nb(Ccur); cp = _cp_nb(Ccur); kk = _k_nb(Ccur)
        kf = _iface_nb(kk)
        _assemble_T_nb(N, dr, h, R0, hc, rho, cp, kk, aT, bT, cT, kf)
        _rhs_T_nb(N, dr, h, R0, hc, rho, cp, Tcur, tinf, dT)
        Tcur = _thomas_nb(aT, bT, cT, dT)

        _rhs_C_nb(N, dr, h, R0, km, Ccur, cinf, dC)
        bse = dC.copy()
        Cit = Ccur.copy()
        for it in range(1, maxit + 1):
            Dn = _faceD(Cit, Tcur, d_fac, mode, nsamp)
            _assemble_C_nb(N, dr, h, Dn, R0, km, aC, bC, cC)
            Ctry = _thomas_nb(aC, bC, cC, bse)
            Cnew = np.empty(NC)
            for i in range(NC):
                Cnew[i] = omega * Ctry[i] + (1.0 - omega) * Cit[i]
            num = 0.0; den = 0.0
            for i in range(NC):
                v = abs(Cnew[i] - Cit[i])
                if v > num: num = v
                w = abs(Cnew[i])
                if w > den: den = w
            if den < 1.0: den = 1.0
            rr = num / den
            Cit = Cnew
            inner += 1
            if rr > resmax: resmax = rr
            if rr < tol: break
        Ccur = Cit
    return Tcur, Ccur, inner, resmax


def load_env():
    wb = openpyxl.load_workbook(ATT1, data_only=True)
    ws = wb.active
    rows = [r for r in ws.iter_rows(values_only=True)][1:]
    d = [r for r in rows if r[0] is not None]
    t = np.array([float(r[0]) for r in d])
    T = np.array([float(r[1]) for r in d])
    C = np.array([float(r[2]) for r in d])
    wb.close()
    return q3.env_tables_from_att1(t, T, C)


def run_fixed(args):
    """固定时程运行（用于 Q2 重叠段影响反查）。"""
    ts, Tv, Cv = load_env()
    nh = int(round(args.hours * 3600.0 / DT_OUT))
    T = np.full(N + 1, DEF2['T0'])
    C = np.full(N + 1, DEF2['C0'])
    snaps = {}
    t0 = time.time()
    for n in range(1, nh + 1):
        T, C, _, _ = _step_nb(N, DR, HH, NSUB, R0, HC, KM, T, C, (n - 1) * DT_OUT,
                              ts, Tv, Cv, D_FAC, args.mode, args.nsamp, 0.7, 1e-10, 30)
        if n in (30, 60, 120, 180):
            snaps[n] = C.copy()
    lines = []
    lines.append('=' * 88)
    lines.append('Q3 诊断②b：界面 D 取法对 **Q2 重叠段（3 h）** 的影响  mode=%d' % args.mode)
    lines.append('=' * 88)
    lines.append('  [SCALE] N=%d Δr=%.2f mm  内部步 1/32 s  固定时程 %.1f h'
                 % (N, DR * 1e3, args.hours))
    wb2 = openpyxl.load_workbook(os.path.join(WS, '20_交付包', '09_代码与复现',
                                              'results', 'result2.xlsx'), read_only=True)
    wsc = wb2['水分浓度']
    r2 = [r for r in wsc.iter_rows(values_only=True)]
    t2v = np.array([float(r[0]) for r in r2[1:]])
    A2 = np.array([[float(x) for x in r[1:]] for r in r2[1:]])
    wb2.close()
    for n in (30, 60, 120, 180):
        tt = n * DT_OUT
        k2 = int(np.argmin(np.abs(t2v - tt)))
        lines.append('   t=%6.0f s  C(0)=%.6f  C(R)=%.6f   与 result2 的 C(0) 差 %+.3e'
                     % (tt, snaps[n][0], snaps[n][-1], snaps[n][0] - A2[k2, 0]))
    lines.append('   21 输出列（3 h 末）：' + ' '.join('%.4f' % snaps[180][i] for i in range(0, 81, 4)))
    lines.append('   墙钟 %.1f s' % (time.time() - t0))
    txt = '\n'.join(lines)
    print(txt)
    with open(LOG, 'a', encoding='utf-8') as f:
        f.write(txt + '\n\n')
    print('已追加：%s' % LOG)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--mode', type=int, required=True, choices=(0, 1, 2))
    ap.add_argument('--nsamp', type=int, default=8)
    ap.add_argument('--max-hours', type=float, default=120.0)
    ap.add_argument('--hours', type=float, default=0.0,
                    help='>0 时改为固定时程（用于 Q2 重叠段影响反查），不触发事件驱动')
    ap.add_argument('--go', action='store_true')
    args = ap.parse_args()
    if not (args.go or os.environ.get('ALLOW_RUN') == '1'):
        print('[SCALE] q3_diag_iface 规模声明：')
        print('        mode=%d（0 调和／1 算术／2 沿C积分 nsamp=%d）  N=%d  Δr=%.2f mm  熔断 %.0f h'
              % (args.mode, args.nsamp, N, DR * 1e3, args.max_hours))
        print('        --hours>0 时跑固定时程（Q2 反查）；否则事件驱动。')
        print('        预计 ≈160–400 s；**仅作根因判别，不产交付物**。加 --go 执行。')
        return

    if args.hours > 0:
        return run_fixed(args)

    ts, Tv, Cv = load_env()
    T = np.full(N + 1, DEF2['T0'])
    C = np.full(N + 1, DEF2['C0'])
    nsteps_max = int(round(args.max_hours * 3600.0 / DT_OUT))
    c24 = None
    t_end = None
    t0 = time.time()
    for n in range(1, nsteps_max + 1):
        t_start = (n - 1) * DT_OUT
        T_prev, C_prev = T.copy(), C.copy()
        T, C, _, _ = _step_nb(N, DR, HH, NSUB, R0, HC, KM, T, C, t_start,
                              ts, Tv, Cv, D_FAC, args.mode, args.nsamp, 0.7, 1e-10, 30)
        if n == 1440:
            c24 = C.copy()
        if float(np.max(C)) < C_CRIT:
            lo, hi = 1, NSUB
            while lo < hi:
                mid = (lo + hi) // 2
                _, Ca, _, _ = _step_nb(N, DR, HH, mid, R0, HC, KM, T_prev, C_prev, t_start,
                                       ts, Tv, Cv, D_FAC, args.mode, args.nsamp, 0.7, 1e-10, 30)
                if float(np.max(Ca)) < C_CRIT:
                    hi = mid
                else:
                    lo = mid + 1
            t_end = t_start + lo * HH
            _, C, _, _ = _step_nb(N, DR, HH, lo, R0, HC, KM, T_prev, C_prev, t_start,
                                  ts, Tv, Cv, D_FAC, args.mode, args.nsamp, 0.7, 1e-10, 30)
            break
    wall = time.time() - t0
    lines = []
    lines.append('=' * 88)
    lines.append('Q3 诊断②：界面 D 取法对照  mode=%d（%s）'
                 % (args.mode, {0: '调和平均/现行', 1: '算术平均', 2: '沿C积分平均 nsamp=%d' % args.nsamp}[args.mode]))
    lines.append('=' * 88)
    lines.append('  [SCALE] N=%d Δr=%.2f mm  内部步 1/32 s  判据 C<%.2f  熔断 %.0f h'
                 % (N, DR * 1e3, C_CRIT, args.max_hours))
    if t_end is not None:
        lines.append('  ★ t_end = %.3f s = %.4f h（相对现行 87.4933 h：%+.4f h / %+.2f%%）'
                     % (t_end, t_end / 3600.0, t_end / 3600.0 - 87.4933,
                        100.0 * (t_end / 3600.0 - 87.4933) / 87.4933))
    else:
        lines.append('  ⚠ 未达标（%.0f h 内）末端 C(0)=%.6f' % (args.max_hours, C[0]))
    lines.append('  末端 C(0)=%.6f  C(R)=%.6f' % (C[0], C[-1]))
    if c24 is not None:
        cols = [0, 20, 40, 60, 80]
        lines.append('  24 h 五列（0/0.5/1.0/1.5/2.0 cm）：' + '  '.join('%.6f' % c24[i] for i in cols))
    lines.append('  墙钟 %.1f s' % wall)
    txt = '\n'.join(lines)
    print(txt)
    with open(LOG, 'a', encoding='utf-8') as f:
        f.write(txt + '\n\n')
    print('已追加：%s' % LOG)


if __name__ == '__main__':
    main()
