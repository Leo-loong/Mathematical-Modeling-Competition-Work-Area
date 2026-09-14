# -*- coding: utf-8 -*-
"""
A 题 Q2 主力求解 · ★「优化前」基线版（前台 · 供用户对照运行）
======================================================
**用途**：用**未做 B/C 优化的原始核**跑一遍完整主力求解，产出独立的结果文件，
与优化版结果做差异核对。

与 `q2_solver.py` 的关系
----------------------
  · 数值设置**完全相同**：Δr=0.25 mm（N=80）、输出步长 1 s、内部步长 1/32 s、
    10800 步（0–3 h）、IMEX 格式、6 项程序化断言。
  · **唯一区别**：求解核为 `q2_core_orig`（scipy `solve_banded` 默认参数、
    装配每次新建数组、无 JIT）⟹ **较慢，预计 5–6 分钟**。
  · **输出隔离**：结果写入 `result2_orig.xlsx`（在**本脚本同目录**），
    **不会覆盖**交付包内的 `result2.xlsx`。

前台模式
--------
  带进度显示（每 1800 步打印一次），结束后**等待按键**再退出。

运行
----
  在本目录下：`python q2_solver_orig.py`
"""
import sys
import os
import time

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

# 线程环境变量须早于 numpy 导入
for _v in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS',
           'NUMEXPR_NUM_THREADS', 'VECLIB_MAXIMUM_THREADS'):
    os.environ.setdefault(_v, '1')

import numpy as np                                              # noqa: E402
from openpyxl import Workbook                                   # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from q2_core_orig import run_q2, DEF2                            # noqa: E402

LOG = os.path.join(HERE, 'q2_run_orig_log.txt')
RES_OUT = os.path.join(HERE, 'result2_orig.xlsx')                # ★ 与优化版隔离

DR = 2.5e-4        # 空间步长（m）
N = 80             # 区间数（节点 81）
DT_OUT = 1.0       # 输出步长（s）
NSTEP = 10800      # 3 h
NSUB = 32          # 内部子步数（内部步长 1/32 s）
COLS = np.arange(0, N + 1, 4)
SNAP = (1800, 3600, 5400, 7200, 9000, 10800)
T_EXT, C_EXT = 50.00, 0.04999


def _find_root(p, marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, marker)):
            return cur
        par = os.path.dirname(cur)
        if par == cur:
            break
        cur = par
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = _find_root(HERE)
ATT1 = os.path.join(ROOT, '10_赛题', 'A题', '附件', '附件1.xlsx')

_fh = open(LOG, 'w', encoding='utf-8')


def say(s=''):
    print(s, flush=True)
    _fh.write(s + '\n')
    _fh.flush()


def load_att1():
    from openpyxl import load_workbook
    wb = load_workbook(ATT1, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr = rows[0]
    data = [r for r in rows[1:] if r[0] is not None]
    return (np.array([float(r[0]) for r in data]),
            np.array([float(r[1]) for r in data]),
            np.array([float(r[2]) for r in data]), hdr)


def main():
    t0 = time.time()
    say('=' * 78)
    say('A 题 · 第二问（Q2）主力求解【优化前基线版 · 前台】')
    say(f'求解核：q2_core_orig（scipy solve_banded 默认参数，无 JIT）')
    say('=' * 78)

    t1, T1, C1, hdr = load_att1()
    say(f'[附件1] 表头={hdr}  点数={len(t1)}  时间 {t1[0]:.0f}–{t1[-1]:.0f} s')
    say(f'        温度 {T1.min():.4f}–{T1.max():.4f} ℃   '
        f'水分浓度 {C1.min():.5f}–{C1.max():.5f} kg/kg')
    say(f'        外推（H6，t>{t1[-1]:.0f} s）：T∞={T_EXT} ℃，C∞={C_EXT} kg/kg')

    def Tenv(t):
        return float(np.interp(t, t1, T1)) if t <= t1[-1] else T_EXT

    def Cenv(t):
        return float(np.interp(t, t1, C1)) if t <= t1[-1] else C_EXT

    say(f'\n[求解] Δr={DR*1e3:.2f} mm（N={N}区间/{N+1}节点）  输出步长={DT_OUT:g} s  '
        f'内部步长=1/{NSUB} s  格式=IMEX')
    say(f'       共 {NSTEP} 个输出步（基线版较慢，请耐心等待）…')
    res, tableT, tableC = run_q2(N=N, dr=DR, dt_out=DT_OUT, nsteps=NSTEP,
                                 Tenv_fn=Tenv, Cenv_fn=Cenv, n_sub=NSUB,
                                 mode='coupled', cols=COLS, snap_at=SNAP,
                                 progress=1800)
    wall = time.time() - t0
    st = res['stats']
    say(f'[完成] 耗时 {wall:.1f} s  内层均 {st["inner_tot"]/max(1,st["nsub"]):.1f} 次/步  '
        f'内层最大残差 {st["resC_max"]:.2e}  异常标记 {st["flags"]}')
    T_end, C_end = res['T_end'], res['C_end']
    say(f'[终态] T(0)={T_end[0]:.6f} ℃  T(R)={T_end[-1]:.6f} ℃  '
        f'C(0)={C_end[0]:.6f}  C(R)={C_end[-1]:.6f}  max_r C={C_end.max():.6f}')

    wb = Workbook(write_only=True)
    for name, tab in (('温度', tableT), ('水分浓度', tableC)):
        ws = wb.create_sheet(name)
        ws.append(['时间/s'] + [f'{c*DR*100:.1f}' for c in COLS])
        for i in range(NSTEP):
            ws.append([i + 1] + [round(float(v), 4) for v in tab[i]])
    wb.save(RES_OUT)
    say(f'[输出] {RES_OUT}  （{os.path.getsize(RES_OUT)/1048576:.2f} MB）')

    say('\n[断言]')
    T_UB = float(np.max(T1)) + 0.01
    T_LB = min(float(np.min(T1)), DEF2['T0']) - 0.01
    chk = []
    chk.append((f'形状 {NSTEP}×{len(COLS)+1}',
                tableT.shape == (NSTEP, len(COLS)) and tableC.shape == (NSTEP, len(COLS))))
    chk.append((f'温度不超环境极值 {T_UB:.4f} ℃', bool(np.nanmax(tableT) <= T_UB)))
    chk.append((f'温度不低于下界 {T_LB:.4f} ℃', bool(np.nanmin(tableT) >= T_LB)))
    chk.append(('含水率非负', bool(np.nanmin(tableC) >= 0.0)))
    chk.append(('含水率单调不增', bool(np.all(np.diff(tableC, axis=0) <= 1e-9))))
    chk.append((f'含水率不超初值 {DEF2["C0"]}',
                bool(np.nanmax(tableC) <= DEF2['C0'] + 1e-4)))
    for name, ok in chk:
        say(f'   {"PASS" if ok else "FAIL"}  {name}')

    say('\n[表3/表4 取用时刻的关键值]  （0.5–3 h）')
    say(f'{"t/h":>6} {"T(0)":>12} {"T(R)":>12} {"C(0)":>14} {"C(R)":>14} {"maxC":>14}')
    for n in SNAP:
        Tv, Cv = res['T_snap'][n], res['C_snap'][n]
        say(f'{n/3600:6.1f} {Tv[0]:12.6f} {Tv[-1]:12.6f} '
            f'{Cv[0]:14.8f} {Cv[-1]:14.8f} {Cv.max():14.8f}')

    say(f'\n总耗时 {wall:.1f} s')
    say(f'日志：{LOG}')
    say(f'结果：{RES_OUT}')
    say('\n★ 请把本窗口结果告知助手，助手将与优化版（result2.xlsx）逐点核对。')


if __name__ == '__main__':
    main()
    try:
        _fh.close()
    except Exception:
        pass
    try:
        input('\n按 Enter 键退出...')
    except Exception:
        pass
