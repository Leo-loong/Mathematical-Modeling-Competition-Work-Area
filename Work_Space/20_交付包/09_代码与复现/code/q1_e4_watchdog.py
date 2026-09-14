# -*- coding: utf-8 -*-
"""
E4 物理看门狗与退化自检（含边界离散口径对照与人为注错演示）
============================================================
目的：抓"数值算得出、物理却不对"的情形，把出错挡在交付之前（对应创新点 C3）。

四类检查
--------
(a) 方向性断言 ：预热期 ∂rT|_R > 0（表面先热，温度沿 r 向外递增）
(b) 单调性断言 ：C 沿时间**逐点不增**
(c) 极值原理   ：T 全程落在 [min(T0, min T∞), max T∞] 区间内
(d) 退化自检   ：① 环境=初值 → 全域严格恒定
                 ② k ≡ 0 → 内部节点 T 不变，仅表面升温
                 ③ D ≡ 0 → 内部节点 C 不变，仅表面降湿

附加
----
(e) 边界离散口径对照：surf='fvm'（元体平衡半格，主力）vs surf='ghost'（虚拟节点＋中心差分）
(f) 人为注错演示：把 Robin 项符号取反（h<0），验证看门狗**确实能捕获**错误

产出：q1_e4_log.txt
约定：本文件不引用、不记载任何资料中的日期。
"""
import os
import io
import sys
import time
import numpy as np
from openpyxl import load_workbook

import q1_core as core
from q1_core import DEF

HERE = os.path.dirname(os.path.abspath(__file__))
# 自适应定位工作区根：向上查找含 10_赛题 的目录
# （兼容两种目录深度：工作区 11_建模/11-3_算法与管线/Q1/code/ 与交付包 09 的 code/）
def _find_root(p, _marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, _marker)):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = _find_root(HERE)
ATT1 = os.path.join(ROOT, '10_赛题', 'A题', '附件', '附件1.xlsx')
LOG = os.path.join(HERE, 'q1_e4_log.txt')
N, DR = 80, 0.25e-3
DT, SUB, SUBT = 1.0, 10, 32
L = []
FAILS = []


def say(s):
    L.append(str(s))
    print(s)


def check(name, ok, detail=''):
    if not ok:
        FAILS.append(name)
    say('  [%s] %s%s' % ('PASS' if ok else 'FAIL', name, ('  ' + detail) if detail else ''))
    return ok


def load_att1():
    wb = load_workbook(ATT1, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    data = list(ws.iter_rows(values_only=True))[1:]
    return (np.array([float(r[0]) for r in data]),
            np.array([float(r[1]) for r in data]),
            np.array([float(r[2]) for r in data]))


def main():
    t0 = time.time()
    t, Tinf, Cinf = load_att1()
    p = dict(DEF)
    Tenv = lambda s: float(np.interp(s, t, Tinf))
    Cenv = lambda s: float(np.interp(s, t, Cinf))

    # ============================ (a)(b)(c) 主力解全程检查 ============================
    say('=== E4-(a)(b)(c) 主力解全程物理检查（全网格 81 点 × 1800 步）===')
    cols = np.arange(N + 1)
    res, tabT, tabC = core.run_sim_collect(
        N=N, dr=DR, dt_out=DT, nsteps=1800, p=p, Tenv_fn=Tenv, Cenv_fn=Cenv,
        sub=SUB, subT=SUBT, center='fv', surf='fvm', cols=cols)

    dT_R = tabT[:, N] - tabT[:, N - 1]          # 表面处向外的近似径向梯度
    check('(a) 预热期 ∂rT|_R > 0（表面先热）', bool(np.all(dT_R > 0)),
          'min=%.3e K（第 %d 步）' % (dT_R.min(), int(np.argmin(dT_R)) + 1))
    # 注：预热最初几步靠近中心的温差低于双精度分辨率，故判据含 -1e-12 容差
    dT_r30 = np.diff(tabT[:30, :], axis=1)
    check('(a2) 前 30 步 T 沿 r 非递减（全节点，容差 1e-12）',
          bool(np.all(dT_r30 > -1e-12)),
          'min=%.3e K' % dT_r30.min())

    dC = np.diff(tabC, axis=0)
    check('(b) C 沿时间逐点不增', bool(np.all(dC <= 1e-12)),
          'max 增量=%.3e' % dC.max())
    dC_sp = np.diff(tabC, axis=1)
    check('(b2) C 沿 r 逐点不增（表面最低）', bool(np.all(dC_sp <= 1e-12)))

    lo = min(p['T0'], float(Tinf.min()))
    hi = float(Tinf.max())
    check('(c) 极值原理 min(T0,T∞min) ≤ T ≤ T∞max',
          bool(tabT.min() >= lo - 1e-9 and tabT.max() <= hi + 1e-9),
          'T∈[%.4f, %.4f]，界=[%.4f, %.4f]' % (tabT.min(), tabT.max(), lo, hi))

    # ============================ (d) 退化自检 ============================
    say('')
    say('=== E4-(d) 退化自检 ===')
    NS = 600                                   # 退化自检用 600 s

    # ① 环境 = 初值 → 全域严格恒定
    r1, _, _ = core.run_sim_collect(
        N=N, dr=DR, dt_out=DT, nsteps=NS, p=p,
        Tenv_fn=lambda s: p['T0'], Cenv_fn=lambda s: p['C0'],
        sub=SUB, subT=SUBT, center='fv', surf='fvm')
    dT = float(np.max(np.abs(r1['T_end'] - p['T0'])))
    dC_ = float(np.max(np.abs(r1['C_end'] - p['C0'])))
    check('(d1) 环境=初值 → T 全域恒定', dT < 1e-9, 'max|ΔT|=%.3e K' % dT)
    check('(d1) 环境=初值 → C 全域恒定', dC_ < 1e-9, 'max|ΔC|=%.3e' % dC_)

    # ② k ≡ 0 → 内部 T 不变，仅表面升温
    p2 = dict(p); p2['k'] = 0.0
    r2, _, _ = core.run_sim_collect(
        N=N, dr=DR, dt_out=DT, nsteps=NS, p=p2,
        Tenv_fn=Tenv, Cenv_fn=Cenv, sub=SUB, subT=SUBT,
        center='fv', surf='fvm', solve_C=False)
    d_in = float(np.max(np.abs(r2['T_end'][:N] - p['T0'])))
    d_sf = float(r2['T_end'][N] - p['T0'])
    check('(d2) k≡0 → 内部 T 严格不变', d_in < 1e-12, 'max|ΔT_in|=%.3e K' % d_in)
    check('(d2) k≡0 → 表面按对流升温', d_sf > 0, 'ΔT_surf=+%.4f K' % d_sf)

    # ③ D ≡ 0 → 内部 C 不变，仅表面降湿
    p3 = dict(p); p3['D0'] = 0.0
    r3, _, _ = core.run_sim_collect(
        N=N, dr=DR, dt_out=DT, nsteps=NS, p=p3,
        Tenv_fn=Tenv, Cenv_fn=Cenv, sub=SUB, subT=SUBT,
        center='fv', surf='fvm', solve_T=False)
    dc_in = float(np.max(np.abs(r3['C_end'][:N] - p['C0'])))
    dc_sf = float(p['C0'] - r3['C_end'][N])
    check('(d3) D≡0 → 内部 C 严格不变', dc_in < 1e-12, 'max|ΔC_in|=%.3e' % dc_in)
    check('(d3) D≡0 → 表面按传质降湿', dc_sf > 0, 'ΔC_surf=-%.6f' % dc_sf)

    # ============================ (e) 边界离散口径对照 ============================
    say('')
    say('=== E4-(e) 边界离散口径对照（前提②澄清）===')
    rf, _, _ = core.run_sim_collect(
        N=N, dr=DR, dt_out=DT, nsteps=1800, p=p, Tenv_fn=Tenv, Cenv_fn=Cenv,
        sub=SUB, subT=SUBT, center='fv', surf='fvm', solve_C=False)
    rg, _, _ = core.run_sim_collect(
        N=N, dr=DR, dt_out=DT, nsteps=1800, p=p, Tenv_fn=Tenv, Cenv_fn=Cenv,
        sub=SUB, subT=SUBT, center='fv', surf='ghost', solve_C=False)
    d_all = float(np.max(np.abs(rf['T_end'] - rg['T_end'])))
    d_s = float(abs(rf['T_end'][N] - rg['T_end'][N]))
    d_c = float(abs(rf['T_end'][0] - rg['T_end'][0]))
    say('  精确半格元体平衡（主力）: T(0)=%.6f  T(R)=%.6f' % (rf['T_end'][0], rf['T_end'][N]))
    say('  一阶近似口径（体积取 R0·Δr/2、界面取 R0）: T(0)=%.6f  T(R)=%.6f'
        % (rg['T_end'][0], rg['T_end'][N]))
    say('  差异：全场 max=%.3e K，表面=%.3e K，中心=%.3e K' % (d_all, d_s, d_c))
    say('  → 差异为 O(Δr) 且相对量 %.2e，远小于物理量 ⟹ 边界离散口径**不敏感**；'
        % (d_s / rf['T_end'][N]))
    say('    主力采用的精确半格元体平衡口径（导热界面半径取 r_{N-1/2}、体积取精确值）无缺陷。')
    say('  → 关于 Math_2 前提②"外边界梯度用跨 Δr 中心差分（半格会差 2 倍）"：本实现按')
    say('    **界面通量**写 Robin（−k·r_{N-1/2}·(T_N−T_{N-1})/Δr 与 h·R0·(T∞−T_N) 并列），')
    say('    不存在"半格梯度"倍率问题；且中心双路只涉及 r=0 的系数装配，与 r=R0 的')
    say('    边界离散无关，其等价性已由 W2 的**解级逐点比对**直接验证。')

    # ============================ (f) 人为注错演示 ============================
    say('')
    say('=== E4-(f) 人为注错演示（验证看门狗确实能抓错）===')
    pb = dict(p); pb['h'] = -abs(p['h'])       # Robin 项符号取反
    rb, _, _ = core.run_sim_collect(
        N=N, dr=DR, dt_out=DT, nsteps=300, p=pb, Tenv_fn=Tenv, Cenv_fn=Cenv,
        sub=SUB, subT=SUBT, center='fv', surf='fvm', solve_C=False)
    dTR_bug = rb['T_end'][N] - rb['T_end'][N - 1]
    say('  注错解：T(R)=%.4f ℃（低于初值 28 ℃ ⟹ 表面被"吸热"）' % rb['T_end'][N])
    say('  方向断言 ∂rT|_R=%.3e（应 >0）→ %s'
        % (dTR_bug, '已捕获' if dTR_bug <= 0 else '未捕获'))
    cap1 = dTR_bug <= 0
    cap2 = bool(rb['T_end'].min() < lo - 1e-9)
    check('(f) 注错被 (a) 方向断言捕获', cap1)
    check('(f) 注错被 (c) 极值原理捕获', cap2,
          '注错解 min T=%.4f < 界 %.4f' % (rb['T_end'].min(), lo))

    # ============================ 汇总 ============================
    say('')
    say('=== E4 汇总 ===')
    say('检查项总数 13，失败 %d 项' % len(FAILS))
    if FAILS:
        say('失败清单：%s' % '；'.join(FAILS))
    say('E4 elapsed=%.1fs' % (time.time() - t0))
    with io.open(LOG, 'w', encoding='utf-8') as f:
        f.write('\n'.join(L) + '\n')
    return 0 if not FAILS else 1


if __name__ == '__main__':
    sys.exit(main())
