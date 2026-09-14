# -*- coding: utf-8 -*-
"""
Q1 诊断脚本（**过程工具，非真值来源**）：
(1) 常数 D 线性解（检验求解器本体）→ 与半无限 Robin 解析估计比较
(2) 非线性 D(C) 解的 Picard 残差序列（检验收敛性）

⚠ **口径警示**：本脚本为**独立自包含实现**，其 C 方程的界面 Df 取
   **距离加权调和平均**（见下方 `assemble_C`），与主力现行口径
   **M6（沿 C 的积分平均）不同**。
   ⟹ 本脚本**仅用于诊断**（P-1／P-2 的定位），其数值**不得用作真值或论文数字**；
   任何 C 侧结论一律以 `q1_core`（主力）与《A_数值口径总表》为准。
"""
import os
import io
import math
import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
# 自适应定位工作区根：向上查找含 10_赛题 的目录
# （兼容两种目录深度：工作区 11_建模/11-3_算法与管线/Q1/code/ 与交付包 09 的 code/）
def _find_root(p, _marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, _marker)):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = _find_root(HERE)
ATT1 = os.path.join(ROOT, '10_赛题', 'A题', '附件', '附件1.xlsx')
LOG = os.path.join(HERE, 'q1_diag_log.txt')

RHO, CP, K, H, KM = 820.0, 2600.0, 0.36, 25.0, 8.0e-7
R0, T0, C0 = 0.02, 28.0, 2.55
DT = 1.0
N = 80
DR = R0 / N
L = []


def say(s):
    L.append(str(s))


def thomas(a, b, c, d):
    n = len(b)
    cp = np.zeros(n)
    dp = np.zeros(n)
    cp[0] = c[0] / b[0]
    dp[0] = d[0] / b[0]
    for i in range(1, n):
        m = b[i] - a[i] * cp[i - 1]
        dp[i] = (d[i] - a[i] * dp[i - 1]) / m
        if i < n - 1:
            cp[i] = c[i] / m
    x = np.zeros(n)
    x[-1] = dp[-1]
    for i in range(n - 2, -1, -1):
        x[i] = dp[i] - cp[i] * x[i + 1]
    return x


def load_cinf():
    wb = load_workbook(ATT1, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    data = list(ws.iter_rows(values_only=True))[1:]
    t = np.array([float(r[0]) for r in data])
    c = np.array([float(r[2]) for r in data])
    return t, c


def assemble_C(D, dc_dr=None):
    """装配 C 方程（常数 D）；dc_dr: 每节点 D（数组）或 None 表示常数 D"""
    r = np.arange(N + 1) * DR
    a = np.zeros(N + 1)
    b = np.zeros(N + 1)
    c = np.zeros(N + 1)
    if dc_dr is None:
        Df = np.full(N, D)
    else:
        Df = np.array([2.0 / (1.0 / dc_dr[i] + 1.0 / dc_dr[i + 1]) for i in range(N)])
    b[0] = 1.0 / DT + 4.0 * Df[0] / DR ** 2
    c[0] = -4.0 * Df[0] / DR ** 2
    for i in range(1, N):
        rl, rr = r[i] - DR / 2.0, r[i] + DR / 2.0
        b[i] = r[i] / DT + (Df[i] * rr + Df[i - 1] * rl) / DR ** 2
        a[i] = -Df[i - 1] * rl / DR ** 2
        c[i] = -Df[i] * rr / DR ** 2
    VN = DR * (R0 - DR / 4.0) / 2.0
    b[N] = VN / DT + Df[N - 1] * (R0 - DR / 2.0) / DR + KM * R0
    a[N] = -Df[N - 1] * (R0 - DR / 2.0) / DR
    return a, b, c, VN, r


def main():
    t, cinf_b = load_cinf()

    # ---------- (1) 常数 D 线性解 ----------
    Dc = 7.0e-9 * math.exp(-0.89 / C0)
    say('D(C0) = %.6e m^2/s' % Dc)
    a, b, c, VN, r = assemble_C(Dc)
    C = np.full(N + 1, C0)
    for n in range(1, 1801):
        cinf = float(np.interp(n * DT, t, cinf_b))
        d = np.zeros(N + 1)
        d[0] = C[0] / DT
        for i in range(1, N):
            d[i] = r[i] * C[i] / DT
        d[N] = VN * C[N] / DT + KM * R0 * cinf
        C = thomas(a, b, c, d)
    say('[linear, const D] after 1800 s: C_center=%.4f  C_surface=%.4f' % (C[0], C[-1]))

    beta = KM * math.sqrt(1800.0) / math.sqrt(Dc)
    cana = float(np.interp(1800.0, t, cinf_b)) + (C0 - float(np.interp(1800.0, t, cinf_b))) \
        * math.exp(beta ** 2) * math.erfc(beta)
    say('[analytic semi-inf Robin] beta=%.3f  C_surface(1800s) approx=%.4f' % (beta, cana))

    # ---------- (2) 非线性 Picard 残差（前若干时步） ----------
    say('')
    say('nonlinear Picard residual trace (first 5 steps, max|dC| per iteration):')
    C = np.full(N + 1, C0)
    for n in range(1, 6):
        cinf = float(np.interp(n * DT, t, cinf_b))
        Dcur = 7.0e-9 * np.exp(-0.89 / np.maximum(C, 1e-6))
        res_trace = []
        for it in range(30):
            aa, bb, cc, VN2, rr = assemble_C(None, dc_dr=Dcur)
            d = np.zeros(N + 1)
            d[0] = C[0] / DT
            for i in range(1, N):
                d[i] = rr[i] * C[i] / DT
            d[N] = VN2 * C[N] / DT + KM * R0 * cinf
            Cn = thomas(aa, bb, cc, d)
            res = float(np.max(np.abs(Cn - C)))
            res_trace.append(res)
            C = Cn
            Dn = 7.0e-9 * np.exp(-0.89 / np.maximum(C, 1e-6))
            rel = float(np.max(np.abs(Dn - Dcur)) / np.max(np.abs(Dcur)))
            if rel < 1e-6:
                break
            Dcur = Dn
        say('  n=%d iters=%d  res=%s' % (n, it + 1, ' '.join('%.2e' % x for x in res_trace[:8])))

    say('')
    say('after 5 steps: C_center=%.6f C_surface=%.6f' % (C[0], C[-1]))

    with io.open(LOG, 'w', encoding='utf-8') as f:
        f.write('\n'.join(L) + '\n')
    print('DONE')


if __name__ == '__main__':
    main()
