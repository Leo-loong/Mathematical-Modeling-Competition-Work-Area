# -*- coding: utf-8 -*-
"""
A 题 Q2 · 图表数据导出（F-15 ～ F-20）
======================================================
输入：result2.xlsx（主力结果）、result1.xlsx（Q1，用于 F-19）、
      w2_foreground_log.txt（W2 预试验，用于 F-20）
输出：20_交付包/04_图表包/data/ 下的 CSV（供写作者按规格卡出图）

产出清单
--------
  fig_q2_field_T.csv    F-16 温度场 r–t（10801 行 × 22 列）
  fig_q2_field_C.csv    F-17 水分浓度场 r–t
  fig_q2_curves.csv     F-18 关键点时程（中心/表面）
  fig_q2_props.csv      F-15 物性随含水率的演化（归一化）
  fig_q1q2_overlap.csv  F-19 Q1 与 Q2 重叠段（0–1800 s）对照
  fig_q2_gci.csv        F-20 收敛性（空间/时间）
"""
import sys
import os
import csv
import re

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np                                                    # noqa: E402
from openpyxl import load_workbook                                    # noqa: E402
from q2_core import rho_of, cp_of, k_of, D_of, DEF2                   # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))


def _find_root(p, marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, marker)):
            return cur
        par = os.path.dirname(cur)
        if par == cur:
            break
        cur = par
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = _find_root(HERE)
RESDIR = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'results')
DATA = os.path.join(ROOT, '20_交付包', '04_图表包', 'data')
DR = 2.5e-4
COLS = list(range(0, 81, 4))          # 0,0.1,…,2.0 cm
C0, T0 = DEF2['C0'], DEF2['T0']


def load_sheet(path, sheet):
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    rows = [r for r in it if r[0] is not None]
    return hdr, rows


def write_csv(name, header, rows):
    p = os.path.join(DATA, name)
    with open(p, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f'  [OK] {name}  {len(rows)} 行')


def f16_f17_f18():
    for sheet, name in (('温度', 'fig_q2_field_T.csv'), ('水分浓度', 'fig_q2_field_C.csv')):
        hdr, rows = load_sheet(os.path.join(RESDIR, 'result2.xlsx'), sheet)
        write_csv(name, ['time_s'] + [f'{c*DR*100:.1f}' for c in COLS],
                  [[int(r[0])] + [float(v) for v in r[1:]] for r in rows])
    # 关键点时程
    ht, rt = load_sheet(os.path.join(RESDIR, 'result2.xlsx'), '温度')
    hc, rc = load_sheet(os.path.join(RESDIR, 'result2.xlsx'), '水分浓度')
    rows = []
    for a, b in zip(rt, rc):
        rows.append([int(a[0]), float(a[1]), float(a[-1]), float(b[1]), float(b[-1])])
    write_csv('fig_q2_curves.csv',
              ['time_s', 'T_center', 'T_surface', 'C_center', 'C_surface'], rows)


def f15_props():
    """物性随含水率变化（归一化到 C0、T0 处）"""
    _, rc = load_sheet(os.path.join(RESDIR, 'result2.xlsx'), '水分浓度')
    Cmin = min(min(float(v) for v in r[1:]) for r in rc)
    Cs = np.linspace(C0, Cmin, 60)
    r0, p0, k0 = float(rho_of(C0)), float(cp_of(C0)), float(k_of(C0))
    D0 = float(D_of(C0, T0))          # 参考点：C0=2.55、T0=28 ℃
    rows = []
    for C in Cs:
        rows.append([round(float(C), 6),
                     round(float(rho_of(C)) / r0, 6),
                     round(float(cp_of(C)) / p0, 6),
                     round(float(k_of(C)) / k0, 6),
                     round(float(D_of(C, T0)) / D0, 6),      # 同温下随 C 变化
                     round(float(D_of(C, 50.0)) / D0, 6)])   # 50 ℃ 下（显温度效应）
    write_csv('fig_q2_props.csv',
              ['C', 'rho_rel', 'cp_rel', 'k_rel', 'D_rel_T28', 'D_rel_T50'], rows)


def f19_overlap():
    """Q1 与 Q2 在 0–1800 s 的对照"""
    _, rt2 = load_sheet(os.path.join(RESDIR, 'result2.xlsx'), '温度')
    _, rc2 = load_sheet(os.path.join(RESDIR, 'result2.xlsx'), '水分浓度')
    _, rt1 = load_sheet(os.path.join(RESDIR, 'result1.xlsx'), '温度')
    _, rc1 = load_sheet(os.path.join(RESDIR, 'result1.xlsx'), '水分浓度')
    n = min(1800, len(rt1), len(rt2))
    rows = []
    for i in range(n):
        rows.append([int(rt2[i][0]),
                     float(rt1[i][1]), float(rt2[i][1]),
                     float(rt1[i][-1]), float(rt2[i][-1]),
                     float(rc1[i][-1]), float(rc2[i][-1])])
    write_csv('fig_q1q2_overlap.csv',
              ['t', 'T_center_q1', 'T_center_q2', 'T_surface_q1', 'T_surface_q2',
               'C_surface_q1', 'C_surface_q2'], rows)


def f20_gci():
    """从 W2 预试验日志解析收敛数据"""
    log = os.path.join(HERE, 'w2_foreground_log.txt')
    if not os.path.exists(log):
        print('  [skip] 未找到 w2_foreground_log.txt')
        return
    rows = []
    for line in open(log, encoding='utf-8'):
        m = re.match(r'\s*(空间 dr=[\d.]+mm|时间 h=1/\d+s)\s+T0=\s*([\d.]+)\s+TR=\s*([\d.]+)'
                     r'\s+C0=([\d.]+)\s+CR=([\d.]+)', line)
        if m:
            tag = m.group(1)
            grp = 'spatial' if tag.startswith('空间') else 'time'
            d = float(re.search(r'dr=([\d.]+)', tag).group(1)) if grp == 'spatial' \
                else 1.0 / float(re.search(r'1/(\d+)', tag).group(1))
            rows.append([grp, tag, d, float(m.group(2)), float(m.group(3)),
                         float(m.group(4)), float(m.group(5))])
    write_csv('fig_q2_gci.csv',
              ['group', 'case', 'delta', 'T0', 'TR', 'C0', 'CR'], rows)


if __name__ == '__main__':
    os.makedirs(DATA, exist_ok=True)
    print('Q2 图表数据导出 →', DATA)
    f16_f17_f18()
    f15_props()
    f19_overlap()
    f20_gci()
    print('完成。')
