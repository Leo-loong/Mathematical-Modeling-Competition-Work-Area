# -*- coding: utf-8 -*-
"""
Q1 图数据导出（供 30_图表/03_图数据准备）
=========================================
从 `20_交付包/09_代码与复现/results/result1.xlsx` 导出三份图数据 CSV：
  · fig_q1_field_T.csv  —— 温度场（time_s × 21 个径向位置），供 F-08 热力图
  · fig_q1_field_C.csv  —— 水分浓度场，供 F-09 热力图
  · fig_q1_curves.csv   —— 中心/表面 温度与含水率时程，供 F-10 曲线

约定：本文件不引用、不记载任何资料中的日期。
"""
import os
import io
import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
# 自适应定位工作区根：向上查找含 10_赛题 的目录
# （兼容两种目录深度：工作区 11_建模/11-3_算法与管线/Q1/code/ 与交付包 09 的 code/）
def _find_root(p, _marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, _marker)):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = _find_root(HERE)
RES1 = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'results', 'result1.xlsx')
OUT = os.path.join(ROOT, '30_图表', '03_图数据准备')


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    hdr = [float(x) for x in rows[0][1:]]
    t = np.array([float(r[0]) for r in rows[1:]])
    A = np.array([[float(v) for v in r[1:]] for r in rows[1:]])
    return hdr, t, A


def main():
    wb = load_workbook(RES1, data_only=True, read_only=True)
    os.makedirs(OUT, exist_ok=True)
    hT, t, AT = read_sheet(wb['温度'])
    hC, _, AC = read_sheet(wb['水分浓度'])
    assert np.allclose(hT, hC), '两表距离列不一致'

    for name, A in (('fig_q1_field_T.csv', AT), ('fig_q1_field_C.csv', AC)):
        with io.open(os.path.join(OUT, name), 'w', encoding='utf-8') as f:
            f.write('time_s,' + ','.join('%.1f' % x for x in hT) + '\n')
            for i in range(len(t)):
                f.write('%.0f,' % t[i] + ','.join('%.4f' % v for v in A[i]) + '\n')

    with io.open(os.path.join(OUT, 'fig_q1_curves.csv'), 'w', encoding='utf-8') as f:
        f.write('time_s,T_center,T_surface,C_center,C_surface\n')
        for i in range(len(t)):
            f.write('%.0f,%.4f,%.4f,%.4f,%.4f\n'
                    % (t[i], AT[i, 0], AT[i, -1], AC[i, 0], AC[i, -1]))

    print('OK  dir=%s  rows=%d  cols=%d' % (OUT, len(t), len(hT)))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
