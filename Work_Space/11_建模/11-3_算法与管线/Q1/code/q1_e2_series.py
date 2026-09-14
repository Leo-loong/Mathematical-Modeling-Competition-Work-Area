# -*- coding: utf-8 -*-
"""
E2 / E2b 圆柱非稳态导热级数解对拍（求解器本体的独立验证）· v2
==============================================================
**v2 变更（scipy 装机后）**
  · 贝塞尔函数改用 `scipy.special.j0/j1`（**全参量域精确**），取代原自研幂级数
    （原实现在参量 x≳45 时中间项达 ~1e17、灾难性抵消失效）；
  · 特征根改用 `scipy.optimize.brentq`，根数由 20 提升至 **200**（特征根上限 ~630）；
  · 因此**取消原"级数适用性"的人为限制**：温度场与含水率场**全部时刻**均可用作对拍判据；
  · 原自研幂级数**保留为安全区一致性核对**（互证两种实现），不再参与生产计算；
  · 保留截断判据 exp(-μmax²Fo) 的**逐点打印**，作为"级数确实收敛"的显式证据。

**E2（常边界）**：常物性 + 常环境 → 与圆柱级数解逐点对拍
**E2b（分段常边界）**：真实时变边界用阶梯近似，数值解与解析解用**同一边界输入**

产出：q1_e2_log.txt；30_图表/03_图数据准备/fig_series_check.csv
约定：本文件不引用、不记载任何资料中的日期。
"""
import os
import io
import sys
import time
import numpy as np
from scipy.special import j0, j1
from scipy.optimize import brentq
from openpyxl import load_workbook

import q1_core as core
from q1_core import DEF

HERE = os.path.dirname(os.path.abspath(__file__))
# 自适应定位工作区根：向上查找含 10_赛题 的目录
# （兼容两种目录深度：工作区 11_建模/11-3_算法与管线/Q1/code/ 与交付包 09 的 code/）
def _find_root(p, _marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, _marker)):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = _find_root(HERE)
ATT1 = os.path.join(ROOT, '10_赛题', 'A题', '附件', '附件1.xlsx')
LOG = os.path.join(HERE, 'q1_e2_log.txt')
CSVDIR = os.path.join(ROOT, '30_图表', '03_图数据准备')
CSV = os.path.join(CSVDIR, 'fig_series_check.csv')
NSTEP = 1800
DT = 1.0
SNAP = (100, 300, 600, 900, 1200, 1500, 1800)
L = []


def say(s):
    L.append(str(s))
    print(s)


# ------------------------------------------ 自研幂级数（仅作安全区一致性核对）
def _j0_series(x):
    x = np.asarray(x, dtype=float)
    half2 = (x / 2.0) ** 2
    term = np.ones_like(x)
    s = np.ones_like(x)
    for k in range(1, 300):
        term = term * (-half2) / (k * k)
        s = s + term
        if np.max(np.abs(term)) < 1e-18:
            break
    return s


def _j1_series(x):
    x = np.asarray(x, dtype=float)
    half = x / 2.0
    half2 = half * half
    term = half.copy() if isinstance(half, np.ndarray) else np.array(half)
    s = term.copy()
    for k in range(1, 300):
        term = term * (-half2) / (k * (k + 1))
        s = s + term
        if np.max(np.abs(term)) < 1e-18:
            break
    return s


def find_roots(Bi, nmax=200, hi=1200.0, nscan=240001):
    """特征方程 μJ1(μ) - Bi·J0(μ) = 0 的正根（scipy brentq）"""
    xs = np.linspace(1e-8, hi, nscan)
    fs = xs * j1(xs) - Bi * j0(xs)
    idx = np.where(np.sign(fs[:-1]) != np.sign(fs[1:]))[0]
    roots = []
    for i in idx:
        a, b = xs[i], xs[i + 1]
        try:
            r = brentq(lambda m: m * j1(m) - Bi * j0(m), a, b, xtol=1e-14, rtol=8.9e-16)
        except Exception:
            continue
        roots.append(float(r))
        if len(roots) >= nmax:
            break
    return np.array(roots)


def series_coeffs(Bi, nmax=200, hi=1200.0):
    mu = find_roots(Bi, nmax, hi)
    J0, J1 = j0(mu), j1(mu)
    A = 2.0 * J1 / (mu * (J0 ** 2 + J1 ** 2))
    return mu, A


def theta(mu, A, rho, Fo):
    """无量纲温度 θ=(T-T∞)/(T0-T∞)"""
    rho = np.asarray(rho, dtype=float)
    s = np.zeros_like(rho)
    for m, a in zip(mu, A):
        s = s + a * j0(m * rho) * np.exp(-m * m * Fo)
    return s


def psi(mu, A, rho, Fo):
    """单位阶跃响应（T∞ 由 0 跳到 1、初值 0 的解）"""
    return 1.0 - theta(mu, A, rho, Fo)


def load_att1():
    wb = load_workbook(ATT1, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    data = list(ws.iter_rows(values_only=True))[1:]
    return (np.array([float(r[0]) for r in data]),
            np.array([float(r[1]) for r in data]),
            np.array([float(r[2]) for r in data]))


def main():
    t0 = time.time()
    t, Tinf, Cinf = load_att1()
    p = dict(DEF)
    R0 = p['R0']
    alpha = p['k'] / (p['rho'] * p['cp'])
    Dc = p['D0'] * np.exp(-p['bD'] / p['C0'])
    Bi_T = p['h'] * R0 / p['k']
    Bi_C = p['km'] * R0 / Dc

    say('=== E2 v2 参数（scipy %s）===' % __import__('scipy').__version__)
    say('alpha=%.6e m^2/s   D(C0)=%.6e m^2/s   Bi_T=%.5f   Bi_m=%.5f'
        % (alpha, Dc, Bi_T, Bi_C))

    # ---- 与自研幂级数的安全区一致性核对 ----
    xs = np.linspace(0.1, 40.0, 40)
    e0 = float(np.max(np.abs(j0(xs) - _j0_series(xs))))
    e1 = float(np.max(np.abs(j1(xs) - _j1_series(xs))))
    say('自研幂级数 vs scipy（x≤40）：max|ΔJ0|=%.3e  max|ΔJ1|=%.3e' % (e0, e1))
    xs2 = np.linspace(0.1, 40.0, 400)
    d0 = np.abs(j0(xs2) - _j0_series(xs2))
    hit = np.where(d0 > 1e-10)[0]
    if len(hit):
        say('自研 J0 首次偏离 scipy 超过 1e-10 的位置：x≈%.2f（此前基本一致）' % xs2[hit[0]])
    say('⟹ **实测结论（修正原判断）**：自研幂级数并非"x≤40 就安全"——')
    say('   · 低阶根（μ≲14，占温度场主导）两实现一致 ⟹ T 场对拍结果不受影响；')
    say('   · 高阶根处自研失效 ⟹ **原版 E2 的级数尾项不准**；含水率场因 Fo 小、')
    say('     恰好依赖高阶项，故其早期解析值存在偏差；本版改用 scipy 后已修正。')
    say('   本版起生产计算一律使用 scipy；自研实现仅作安全区互证，不再参与计算。')
    say('')

    muT, AT = series_coeffs(Bi_T)
    muC, AC = series_coeffs(Bi_C)
    say('T 特征根：共 %d 个，前 5 = %s，μmax=%.2f'
        % (len(muT), ' '.join('%.5f' % x for x in muT[:5]), muT[-1]))
    say('C 特征根：共 %d 个，前 5 = %s，μmax=%.2f'
        % (len(muC), ' '.join('%.5f' % x for x in muC[:5]), muC[-1]))
    say('')

    # ---------------------------------------------------------------- E2a
    say('=== E2a 常边界对拍（T∞=50.0 ℃, C∞=0.05；dr = 1 / 0.5 / 0.25 mm, dt=1 s）===')
    Tinf0, Cinf0 = 50.0, 0.05
    rows = []
    for (N, dr) in [(20, 1.0e-3), (40, 0.5e-3), (80, 0.25e-3)]:
        cols = np.arange(0, N + 1, max(1, N // 20))
        res, _, _ = core.run_sim_collect(
            N=N, dr=dr, dt_out=DT, nsteps=NSTEP, p=p,
            Tenv_fn=lambda s: Tinf0, Cenv_fn=lambda s: Cinf0,
            sub=4, subT=1, center='fv', surf='fvm', dmode='const',
            cols=cols, snap_at=set(SNAP))
        say('--- N=%d (dr=%.2f mm) ---' % (N, dr * 1e3))
        for ts in SNAP:
            Fo = alpha * ts / R0 ** 2
            FoC = Dc * ts / R0 ** 2
            tailT = float(np.exp(-muT[-1] ** 2 * Fo))
            tailC = float(np.exp(-muC[-1] ** 2 * FoC))
            Ta_c = Tinf0 + (p['T0'] - Tinf0) * float(theta(muT, AT, 0.0, Fo))
            Ta_s = Tinf0 + (p['T0'] - Tinf0) * float(theta(muT, AT, 1.0, Fo))
            Ca_c = Cinf0 + (p['C0'] - Cinf0) * float(theta(muC, AC, 0.0, FoC))
            Ca_s = Cinf0 + (p['C0'] - Cinf0) * float(theta(muC, AC, 1.0, FoC))
            Tn_c = float(res['T_snap'][ts][0]); Tn_s = float(res['T_snap'][ts][-1])
            Cn_c = float(res['C_snap'][ts][0]); Cn_s = float(res['C_snap'][ts][-1])
            say(' t=%4d [tail T=%.1e C=%.1e]  T0: %.6f/%.6f d=%+.2e | TR: %.6f/%.6f d=%+.2e'
                % (ts, tailT, tailC, Tn_c, Ta_c, Tn_c - Ta_c, Tn_s, Ta_s, Tn_s - Ta_s))
            say('          C0: %.6f/%.6f d=%+.2e | CR: %.6f/%.6f d=%+.2e'
                % (Cn_c, Ca_c, Cn_c - Ca_c, Cn_s, Ca_s, Cn_s - Ca_s))
            rows.append(('E2a_dr%.2fmm' % (dr * 1e3), ts, Tn_c, Ta_c, Tn_s, Ta_s, Cn_s, Ca_s))
        Fo = alpha * 1800.0 / R0 ** 2
        FoC = Dc * 1800.0 / R0 ** 2
        say('  终态偏差: |ΔT(0)|=%.3e  |ΔT(R)|=%.3e  |ΔC(R)|=%.3e'
            % (abs(float(res['T_snap'][1800][0]) - (Tinf0 + (p['T0'] - Tinf0) * float(theta(muT, AT, 0.0, Fo)))),
               abs(float(res['T_snap'][1800][-1]) - (Tinf0 + (p['T0'] - Tinf0) * float(theta(muT, AT, 1.0, Fo)))),
               abs(float(res['C_snap'][1800][-1]) - (Cinf0 + (p['C0'] - Cinf0) * float(theta(muC, AC, 1.0, FoC))))))

    # ---------------------------------------------------------------- E2b
    say('')
    say('=== E2b 分段常边界对拍（阶梯：每 60 s 取附件1 该段值；同一阶梯输入两侧）===')
    say('  ⚠ 口径说明（本版）：温度已改为「时间精确推进」，其步内假设边界**线性**；')
    say('     而阶梯函数在节点处**跳变**，数值侧只能把它线性化到 1 s 步内 ⟹ 本条残差 =')
    say('     「阶梯不连续性被 1 s 线性化」的差异，**不再反映求解器误差**。')
    say('     真实连续时变边界的对拍已由 **T1-1 Duhamel 半解析**承担（见 innov/logs/q1_inv_T1_duhamel.log）。')
    seg = 60.0

    def step_fn(s):
        k = int(s // seg)
        k = min(max(k, 0), len(Tinf) - 1)
        return float(Tinf[k])

    N, dr = 80, 0.25e-3
    cols = np.arange(0, N + 1, max(1, N // 20))
    res, _, _ = core.run_sim_collect(
        N=N, dr=dr, dt_out=DT, nsteps=NSTEP, p=p,
        Tenv_fn=step_fn, Cenv_fn=lambda s: Cinf0, sub=4, subT=1,
        center='fv', surf='fvm', dmode='const', cols=cols, snap_at=set(SNAP))
    lv = [float(step_fn(i * seg)) for i in range(int(1800 // seg) + 1)]
    say('阶梯段值（前 5）: %s' % ' '.join('%.3f' % v for v in lv[:5]))
    for ts in SNAP:
        Tc_an = p['T0']; Ts_an = p['T0']
        for j in range(1, len(lv)):
            tj = j * seg
            if tj >= ts:
                break
            d = lv[j] - lv[j - 1]
            if d == 0:
                continue
            Fo = alpha * (ts - tj) / R0 ** 2
            Tc_an += d * float(psi(muT, AT, 0.0, Fo))
            Ts_an += d * float(psi(muT, AT, 1.0, Fo))
        Tc_num = float(res['T_snap'][ts][0]); Ts_num = float(res['T_snap'][ts][-1])
        say(' t=%4d  T(0): num=%.6f ana=%.6f d=%+.2e | T(R): num=%.6f ana=%.6f d=%+.2e'
            % (ts, Tc_num, Tc_an, Tc_num - Tc_an, Ts_num, Ts_an, Ts_num - Ts_an))
        rows.append(('E2b_step', ts, Tc_num, Tc_an, Ts_num, Ts_an, float('nan'), float('nan')))

    # ---------------------------------------------------------------- CSV
    os.makedirs(CSVDIR, exist_ok=True)
    with io.open(CSV, 'w', encoding='utf-8') as f:
        f.write('case,t,Tc_num,Tc_ana,Ts_num,Ts_ana,Cs_num,Cs_ana\n')
        for c, ts, a, b, c2, d2, e, g in rows:
            f.write('%s,%d,%.8f,%.8f,%.8f,%.8f,%.8f,%.8f\n' % (c, ts, a, b, c2, d2, e, g))
    say('')
    say('CSV written: %s' % CSV)
    say('E2 elapsed=%.1fs' % (time.time() - t0))
    with io.open(LOG, 'w', encoding='utf-8') as f:
        f.write('\n'.join(L) + '\n')
    return 0


if __name__ == '__main__':
    sys.exit(main())
