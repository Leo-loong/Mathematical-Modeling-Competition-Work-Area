# -*- coding: utf-8 -*-
# 用途：Q4 证据链检查·数值复算（只读，不写任何文件）
#   result4 规模与列契约 / 表6 逐格比对 / 末端态 / 空域掩蔽 / Sobol CSV / table6.csv
"""Q4 独立核验（只读）：result4.xlsx 规模/表头/表6 逐格/末端态/空值契约/Sobol CSV。"""
import io
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))


def _find_root(p, _marker="10_赛题", _max=8):
    """向上探测含 10_赛题 的目录作为工作区根（与项目既有脚本同一约定）。"""
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, _marker)):
            return cur
        cur = os.path.dirname(cur)
    raise SystemExit("未找到工作区根")


ROOT = _find_root(HERE)
Q4 = os.path.join(ROOT, "11_建模", "11-3_算法与管线", "Q4")

# 表6 真值（结果分析 §二）
T6 = {
    6: [1.6276, 1.4621, 0.9868, 0.4173],
    12: [0.7021, 0.6248, 0.3930, 0.1598],
    18: [0.3958, 0.3576, 0.2340, 0.0869],
    24: [0.2794, 0.2561, 0.1763, 0.0665],
    30: [0.2232, 0.2066, 0.1477, 0.0591],
    36: [0.1909, 0.1779, 0.1306, 0.0558],
    42: [0.1699, 0.1592, 0.1193, 0.0540],
    48: [0.1552, 0.1460, 0.1110, 0.0529],
}

print("=" * 70)
print("[1] result4.xlsx（工作区真源）")
p = os.path.join(Q4, "results", "result4.xlsx")
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
print("  sheets =", wb.sheetnames)
ws = wb[wb.sheetnames[0]]
print("  max_row x max_col = %d x %d" % (ws.max_row, ws.max_column))
rows = list(ws.iter_rows(values_only=True))
hdr = rows[0]
print("  header[0:4] =", hdr[:4], " ... header[-3:] =", hdr[-3:])
times = [r[0] for r in rows[1:]]
print("  A col: n=%d, first=%s, second=%s, last=%s" % (len(times), times[0], times[1], times[-1]))
tmap = {}
for i, t in enumerate(times, start=2):
    tmap[round(float(t), 3)] = i

# 定位 0 / 0.5 / 1.0 / 表面 列
def col_of(x):
    for j, h in enumerate(hdr):
        if h is None:
            continue
        s = str(h)
        try:
            if abs(float(s) - x) < 1e-9:
                return j
        except ValueError:
            if ("表面" in s) and x < 0:
                return j
    return None

c0, c5, c10, cR = col_of(0.0), col_of(0.5), col_of(1.0), col_of(-1.0)
print("  列索引: 0cm=%s 0.5cm=%s 1.0cm=%s 表面=%s" % (c0, c5, c10, cR))
print("  1.5cm 是否在表头: %s" % (col_of(1.5) is not None))
print("  1.9cm 是否在表头: %s" % (col_of(1.9) is not None))

print("\n[2] 表6 逐格核对（60 s 行直接取样）")
bad = 0
for h, exp in sorted(T6.items()):
    t = h * 3600
    r = rows[tmap[t] - 1]
    got = [r[c0], r[c5], r[c10], r[cR]]
    line = "  %2dh " % h
    for g, e in zip(got, exp):
        gv = float(g)
        ok = abs(round(gv, 4) - e) < 5e-5
        if not ok:
            bad += 1
        line += "%s(%.4f/%.4f) " % ("OK" if ok else "XX", gv, e)
    print(line)
print("  ⟹ 表6 不符格数 = %d / 32" % bad)

print("\n[3] 末端态（最后一行）")
last = rows[-1]
print("  t = %s s" % last[0])
print("  C(0) = %r  (应 0.14999954, <0.15)" % last[c0])
print("  C(R) = %r  (应 0.0526)" % last[cR])
below = sum(1 for x in last[1:22] if x is not None and float(x) < 0.15)
nonnull = sum(1 for x in last[1:22] if x is not None)
print("  末行非空格点数 = %d，其中 <0.15 的 = %d" % (nonnull, below))

print("\n[4] 空值契约（r > R(t) 留空）")
for h in (6, 24, 48):
    t = h * 3600
    r = rows[tmap[t] - 1]
    vals = [(i, r[i]) for i in range(1, 22)]
    nn = [i for i, v in vals if v is not None]
    print("  t=%2dh: 非空列表数=%d，列序 %s" % (h, len(nn), [str(hdr[i]) for i in nn][-3:]))

print("\n[5] 列数口径（题面要求 0.1 cm 步长 + 表面列）")
dist_cols = [hdr[i] for i in range(1, 22)]
print("  21 个距离列 =", dist_cols[:5], "...", dist_cols[-3:])

print("\n[6] Sobol CSV（IN-6）")
for fn in ("fig_q4_INV6_sobol.csv", "fig_q4_INV6_sobol_S2.csv"):
    f = os.path.join(Q4, "innov", "out", fn)
    with io.open(f, "r", encoding="utf-8") as fh:
        txt = fh.read().strip().splitlines()
    print("  --- %s (%d 行)" % (fn, len(txt)))
    for ln in txt:
        print("      " + ln)

print("\n[7] table6.csv 表头与内容")
f = os.path.join(Q4, "results", "table6.csv")
with io.open(f, "r", encoding="utf-8") as fh:
    for ln in fh.read().strip().splitlines():
        print("  " + ln)
print("=" * 70)
