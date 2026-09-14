# -*- coding: utf-8 -*-
# 用途：Q4 证据链检查·半径与计数复算（只读）
#   附件2 R(t) 复算（R(t_dry) 与各时刻半径）/ 交付包脚本计数 / 支撑材料清单编号核查
"""Q4 独立核验（二）：R(t) 附件2 复算、交付包脚本计数、支撑材料清单核对。"""
import io
import os
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))


def _find_root(p, _marker="10_赛题", _max=8):
    """向上探测含 10_赛题 的目录作为工作区根（与项目既有脚本同一约定）。"""
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, _marker)):
            return cur
        cur = os.path.dirname(cur)
    raise SystemExit("未找到工作区根")


ROOT = _find_root(HERE)
Q4 = os.path.join(ROOT, "11_建模", "11-3_算法与管线", "Q4")

print("=" * 70)
print("[1] 附件2 R(t) 复算（R(t_dry) 应为多少？）")
cand = []
base = os.path.join(ROOT, "10_赛题", "A题", "附件")
for dp, dns, fns in os.walk(base):
    for fn in fns:
        if "2" in fn and ("附件" in dp.split(os.sep)[-1] or True):
            cand.append(os.path.join(dp, fn))
print("  附件目录文件：")
for c in cand:
    print("    %s  (%d B)" % (os.path.relpath(c, ROOT), os.path.getsize(c)))

x2 = [c for c in cand if os.path.splitext(c)[1].lower() in (".xlsx", ".xls")]
if x2:
    p = x2[0]
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [r for r in ws.iter_rows(values_only=True) if r and r[0] is not None]
    print("  文件 = %s ; sheet=%s ; 行数=%d" % (os.path.basename(p), wb.sheetnames, len(rows)))
    print("  前 3 行 = %s" % (rows[:3],))
    print("  后 3 行 = %s" % (rows[-3:],))
    data = []
    for r in rows:
        try:
            t, R = float(r[0]), float(r[1])
        except (TypeError, ValueError):
            continue
        data.append((t, R))
    print("  有效点 %d 个；R(0)=%.6f R(end)=%.6f t_end=%.1f" % (len(data), data[0][1], data[-1][1], data[-1][0]))
    # 单位判断
    print("  ⟹ R 单位疑似 %s（若首值≈2.0 则 cm，≈0.02 则 m）" % ("cm" if data[0][1] > 0.5 else "m"))

    def R_at(t):
        if t <= data[0][0]:
            return data[0][1]
        if t >= data[-1][0]:
            return data[-1][1]
        for i in range(1, len(data)):
            if data[i][0] >= t:
                t0, R0 = data[i - 1]
                t1, R1 = data[i]
                w = (t - t0) / (t1 - t0) if t1 > t0 else 0.0
                return R0 + w * (R1 - R0)
        return data[-1][1]

    scale = 100.0 if data[0][1] < 0.5 else 1.0   # -> cm
    print("  表6/末端 R 复算（单位 cm）：")
    for h in (6, 12, 18, 24, 30, 36, 42, 48):
        print("    t=%2dh  R=%.4f cm" % (h, R_at(h * 3600) * scale))
    td = 182348.109
    print("    t=t_dry=%.3f s  R=%.6f cm" % (td, R_at(td) * scale))
    print("    R 首次到达最小值 %.4f cm 的时刻 = %.1f s (%.2f h)"
          % (data[-1][1] * scale, next(t for t, R in data if R <= data[-1][1] + 1e-12),
             next(t for t, R in data if R <= data[-1][1] + 1e-12) / 3600.0))

print()
print("=" * 70)
print("[2] 交付包脚本计数（实测）")
CODE = os.path.join(ROOT, "20_交付包", "09_代码与复现", "code")
g = {}
tot = 0
for fn in sorted(os.listdir(CODE)):
    if not fn.endswith(".py"):
        continue
    tot += 1
    m = re.match(r"^(q\d|dq)", fn)
    k = m.group(1) if m else "其他"
    g.setdefault(k, []).append(fn)
print("  code/ 下 .py 总数 = %d" % tot)
for k in sorted(g):
    print("    %-6s = %d" % (k, len(g[k])))
inn = os.path.join(CODE, "innov")
if os.path.isdir(inn):
    py = [f for f in os.listdir(inn) if f.endswith(".py")]
    print("  code/innov/ .py = %d : %s" % (len(py), sorted(py)))

print()
print("[3] 支撑材料清单：编号重复与 Q4 条目")
f = os.path.join(ROOT, "20_交付包", "11_支撑材料包", "A_支撑材料文件列表.md")
txt = io.open(f, "r", encoding="utf-8").read().splitlines()
seen = {}
for i, ln in enumerate(txt, 1):
    m = re.match(r"^\|\s*(\d+)\s*\|", ln)
    if m:
        n = int(m.group(1))
        seen.setdefault(n, []).append(i)
dup = {k: v for k, v in seen.items() if len(v) > 1}
print("  编号重复：", dict(sorted(dup.items())) if dup else "无")
print("  最大编号 = %s，条目数 = %d" % (max(seen) if seen else None, len(seen)))
q4lines = [(i, ln) for i, ln in enumerate(txt, 1) if "q4_" in ln]
print("  提到 q4_ 的行数 = %d" % len(q4lines))
for i, ln in q4lines[:6]:
    print("    %d: %s" % (i, ln[:110]))
for k in ("85", "84", "18 个", "17 个", "24 个"):
    hit = [i for i, ln in enumerate(txt, 1) if k in ln]
    print("  含 %-6s 的行：%s" % (k, hit[:6]))
print("=" * 70)
