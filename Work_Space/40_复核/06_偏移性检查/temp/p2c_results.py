# -*- coding: utf-8 -*-
"""P2c 第三向：**结果文件**侧核验——口径总表 §十三 登记的"论文显现值"是否真在 result*.xlsx 中。只读。"""
import os, csv, io
from openpyxl import load_workbook

ROOT = r'c:\Users\wang-\Desktop\2026数学建模\Work_Space'
RES = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'results')

TARGETS = [
    ('result2.xlsx', 1.766193, '问题二 3 h 中心含水率（§十三 第 1 条）'),
    ('result2.xlsx', 1.008113, '问题二 3 h 表面含水率（§十三 第 2 条）'),
    ('result3.xlsx', 207113.094, '问题三达标时刻（秒制，§十三 第 3 条）'),
    ('result3.xlsx', 57.5314, '问题三达标时刻（小时制）'),
    ('result4.xlsx', 182348.109, '问题四烘干时刻（秒制，§十三 第 4 条）'),
    ('result4.xlsx', 50.6523, '问题四烘干时刻（小时制）'),
]


def nums_of(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, (int, float)):
                    out.append(float(v))
    return out


cache = {}
print('=== P2c 结果文件侧核验（论文显现值 ↔ result*.xlsx）===')
for fn, val, desc in TARGETS:
    p = os.path.join(RES, fn)
    if not os.path.isfile(p):
        print('   ⚠ 缺件 %s' % fn)
        continue
    if fn not in cache:
        cache[fn] = nums_of(p)
        print('   （%s：读出 %d 个数值单元）' % (fn, len(cache[fn])))
    hit = [x for x in cache[fn] if abs(x - val) <= 5e-4]
    print('   %-14s %-14s → %s  %s' % (fn, val, '✅ 命中' if hit else '❌ **未命中**', desc))

# 与正文侧交叉：正文是否出现同一值
pdf = os.path.join(ROOT, '50_论文', '04_交排版成图方', '交付排版方论文', 'main.pdf')
try:
    import fitz
    d = fitz.open(pdf)
    txt = ''.join(p.get_text() for p in d)
    print()
    print('--- 同一值在正文 PDF 中的出现 ---')
    for _, val, desc in TARGETS:
        for form in ('%.3f' % val, '%.4f' % val, '%.2f' % (val / 3600) if val > 1000 else '%.4f' % val):
            if form in txt:
                print('   %-14s ← 纸面出现「%s」  %s' % (val, form, desc))
                break
        else:
            print('   %-14s ← 纸面**未见**（可能以换算形式呈现）  %s' % (val, desc))
except Exception as e:
    print('（PDF 侧跳过：%s）' % e)
