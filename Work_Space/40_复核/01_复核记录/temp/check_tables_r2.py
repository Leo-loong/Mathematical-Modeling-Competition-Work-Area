# -*- coding: utf-8 -*-
"""第二轮逐格比对：从现行 A_08 解析表1–6 ↔ result*.xlsx（只读）"""
import re, sys, io
from pathlib import Path
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(r"C:\Users\wang-\Desktop\2026数学建模\Work_Space")
md = (ROOT / "50_论文/02_章节稿/A_08_模型建立与求解_基线.md").read_text(encoding="utf-8")
RES = ROOT / "20_交付包/09_代码与复现/results"

def parse_table(title):
    i = md.find("**" + title)
    if i < 0: return None, None
    seg = md[i:]
    rows = []
    for ln in seg.splitlines()[1:]:
        ln = ln.strip()
        if not ln.startswith("|"):
            if rows: break
            continue
        cells = [c.strip().replace("**", "") for c in ln.strip("|").split("|")]
        if set(cells[0]) <= set("-: "):  # 分隔行
            continue
        rows.append(cells)
    hdr = [h for h in rows[0]] if rows else []
    return hdr, rows

def r4(x): return round(float(x), 4)

def load(path, sheet=None):
    wb = load_workbook(path, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    data, first, last = {}, True, None
    for row in ws.iter_rows(values_only=True):
        if first: first = False; continue
        if row[0] is None: continue
        t = round(float(row[0]), 3)
        data[t] = row; last = row
    wb.close(); return data, last

# 结果文件列映射：0-based（0=时间）
CMAP5 = {("0"):1, ("0.5"):6, ("1.0"):11, ("1.5"):16, ("2.0"):21}
CMAP6 = {("0"):1, ("0.5"):6, ("1.0"):11, ("药材表面"):21}

specs = [
    ("表 1", "result1.xlsx", "温度", "s", CMAP5),
    ("表 2", "result1.xlsx", "水分浓度", "s", CMAP5),
    ("表 3", "result2.xlsx", "温度", "s", CMAP5),
    ("表 4", "result2.xlsx", "水分浓度", "s", CMAP5),
    ("表 5", "result3.xlsx", None, "h", CMAP5),
    ("表 6", "result4.xlsx", None, "h", CMAP6),
]
books = {}
for f in ("result1.xlsx","result2.xlsx","result3.xlsx","result4.xlsx"):
    books[f] = load(RES / f)

grand_ok = grand_bad = grand_na = 0
for title, fname, sheet, mode, cmap in specs:
    hdr, rows = parse_table(title)
    data, last = books[fname]
    ok = bad = na = 0; bads = []
    for r in rows:
        tlab = r[0]
        if tlab.startswith("烘干结束") or "烘干结束" in tlab:
            row = last
        else:
            tv = float(tlab)
            tsec = tv if mode == "s" else round(tv * 3600, 3)
            row = data.get(round(tsec, 3)) or data.get(int(tsec) if float(tsec).is_integer() else tsec)
        if row is None:
            bads.append((tlab, "行缺失")); continue
        for j, cell in enumerate(r[1:], start=1):
            colname = hdr[j]
            if cell in ("—", "", "—", None):
                na += 1; continue
            ci = cmap.get(colname) or cmap.get(colname.replace(" ", ""))
            if ci is None:
                bads.append((tlab, f"列'{colname}'无映射")); continue
            xv = row[ci] if ci < len(row) else None
            if xv is None or xv == "":
                bads.append((tlab, f"列{colname} 文件空(论文{cell})")); bad += 1; continue
            if abs(r4(xv) - float(cell)) > 1e-9:
                bads.append((tlab, colname, f"论文{cell} vs 文件{r4(xv)}")); bad += 1
            else:
                ok += 1
    grand_ok += ok; grand_bad += bad; grand_na += na
    print(f"[{title}] 一致 {ok} ｜ 不一致 {bad} ｜ 空/跳过 {na}")
    for b in bads[:10]: print("   ", b)
print(f"\n总计：一致 {grand_ok} ｜ 不一致 {grand_bad} ｜ 空/跳过 {grand_na}")
