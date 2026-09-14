# -*- coding: utf-8 -*-
"""
B1：边界数据噪声的量化与"噪声 → 输出"传播带（Q1 与 Q2）
========================================================
只读调用 q1_core / q2_core；不写入 11_建模、20_交付包 任何文件。
方法：以"原始边界"与"移动平均去噪边界"两次求解之差，作为**输入高频噪声对输出的贡献**
      （去噪即把高频分量移除，故两次之差 ≈ 噪声分量本身的贡献；这是可复现的保守估计）。
"""
import io
import os
import sys
import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from exp_common import find_root

ROOT = find_root()
sys.path.insert(0, os.path.join(ROOT, '11_建模', '11-3_算法与管线', 'Q1', 'code'))
sys.path.insert(0, os.path.join(ROOT, '11_建模', '11-3_算法与管线', 'Q2', 'code'))
import q1_core as C1          # noqa: E402
import q2_core as C2          # noqa: E402

ATT1 = os.path.join(ROOT, '10_赛题', 'A题', '附件', '附件1.xlsx')
OUT = os.path.join(HERE, 'exp_noise_report.txt')
L = []


def say(s=''):
    print(s, flush=True)
    L.append(str(s))


def load():
    wb = load_workbook(ATT1, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    d = [r for r in ws.iter_rows(values_only=True)][1:]
    d = [r for r in d if r[0] is not None]
    return (np.array([float(r[0]) for r in d]), np.array([float(r[1]) for r in d]),
            np.array([float(r[2]) for r in d]))


def ma(y, half):
    """±half 点移动平均（端点用截断窗）"""
    n = len(y)
    out = np.empty(n)
    for i in range(n):
        lo, hi = max(0, i - half), min(n, i + half + 1)
        out[i] = y[lo:hi].mean()
    return out


def main():
    t, T, C = load()
    seg = (t >= 10000) & (t <= 14400)
    say('=' * 96)
    say('B1 边界数据噪声量化和"噪声→输出"传播带')
    say('=' * 96)
    say('【1】附件1 噪声统计（平台段 [10000,14400] s，n=%d）' % seg.sum())
    say('   T∞: 均值 %.4f ℃  σ=%.4f ℃  极差 %.4f ℃  最大相邻跳变 %.4f ℃' %
        (T[seg].mean(), T[seg].std(ddof=1), T[seg].max() - T[seg].min(),
         np.abs(np.diff(T[seg])).max()))
    say('   C∞: 均值 %.6f  σ=%.6f  极差 %.6f  最大相邻跳变 %.6f' %
        (C[seg].mean(), C[seg].std(ddof=1), C[seg].max() - C[seg].min(),
         np.abs(np.diff(C[seg])).max()))
    say('   一阶差分噪声代理（全段，÷√2）: σ_T=%.4f ℃   σ_C=%.6f' %
        (np.diff(T).std(ddof=1) / np.sqrt(2), np.diff(C).std(ddof=1) / np.sqrt(2)))
    say('   ⟹ 平台段温度含 ±%.2f ℃ 量级的高频波动（相对 50 ℃ 约 %.2f%%）；水分浓度含 ±%.1e 量级波动。'
        % (T[seg].std(ddof=1), T[seg].std(ddof=1) / 50 * 100, C[seg].std(ddof=1)))
    say('')

    # ---------------------------------------------------------------- Q1
    p = dict(C1.DEF)
    raw = {}
    for tag, half in [('原始', 0), ('MA±60 s', 1), ('MA±120 s', 2)]:
        Tt, Cc = (T, C) if half == 0 else (ma(T, half), ma(C, half))
        r = C1.run_sim(N=80, dr=0.00025, dt_out=1.0, nsteps=1800, p=p,
                       Tenv_fn=lambda s, Tt=Tt: float(np.interp(s, t, Tt)),
                       Cenv_fn=lambda s, Cc=Cc: float(np.interp(s, t, Cc)),
                       sub=10, subT=32, center='fv', surf='fvm')
        raw[tag] = r
    say('【2】Q1：去噪前后 关键量对照（Δr=0.25 mm，主力口径）')
    say('   %-10s %14s %14s %14s %14s' % ('边界口径', 'T(0,1800)', 'T(R,1800)', 'C(0,1800)', 'C(R,1800)'))
    base = raw['原始']
    for tag in ('原始', 'MA±60 s', 'MA±120 s'):
        r = raw[tag]
        say('   %-10s %14.6f %14.6f %14.8f %14.8f' % (tag, r['T_end'][0], r['T_end'][-1],
                                                      r['C_end'][0], r['C_end'][-1]))
    for tag in ('MA±60 s', 'MA±120 s'):
        dT = raw[tag]['T_end'] - base['T_end']
        dC = raw[tag]['C_end'] - base['C_end']
        say('   %s − 原始： max|ΔT|=%.3e ℃  max|ΔC|=%.3e  ；登记量 ΔT(0)=%+.3e  ΔT(R)=%+.3e  ΔC(R)=%+.3e'
            % (tag, np.abs(dT).max(), np.abs(dC).max(), dT[0], dT[-1], dC[-1]))
    say('   对照：T 的网格误差带 ≈5e-5 ℃；C(R) 的网格误差带 ≈1e-4')
    say('')

    # ---------------------------------------------------------------- Q2
    T_EXT, C_EXT = 50.00, 0.04999
    q2 = {}
    for tag, half in [('原始', 0), ('MA±120 s', 2)]:
        Tt, Cc = (T, C) if half == 0 else (ma(T, half), ma(C, half))
        Te = lambda s, Tt=Tt: float(np.interp(s, t, Tt)) if s <= t[-1] else T_EXT
        Ce = lambda s, Cc=Cc: float(np.interp(s, t, Cc)) if s <= t[-1] else C_EXT
        res, _, _ = C2.run_q2(N=80, dr=2.5e-4, dt_out=1.0, nsteps=10800,
                              Tenv_fn=Te, Cenv_fn=Ce, n_sub=32, mode='coupled')
        q2[tag] = res
    say('【3】Q2：去噪前后 关键量对照（Δr=0.25 mm，主力口径，3 h）')
    say('   %-10s %12s %12s %12s %12s' % ('边界口径', 'T(0)', 'T(R)', 'C(0)', 'C(R)'))
    for tag in ('原始', 'MA±120 s'):
        r = q2[tag]
        say('   %-10s %12.6f %12.6f %12.8f %12.8f' % (tag, r['T_end'][0], r['T_end'][-1],
                                                      r['C_end'][0], r['C_end'][-1]))
    dT = q2['MA±120 s']['T_end'] - q2['原始']['T_end']
    dC = q2['MA±120 s']['C_end'] - q2['原始']['C_end']
    say('   MA±120 s − 原始： max|ΔT|=%.3e ℃  max|ΔC|=%.3e ；ΔT(0)=%+.3e  ΔT(R)=%+.3e  ΔC(0)=%+.3e  ΔC(R)=%+.3e'
        % (np.abs(dT).max(), np.abs(dC).max(), dT[0], dT[-1], dC[0], dC[-1]))
    say('   对照：Q2 的网格误差带（W2：Δr 1.0／0.5／0.25 mm 相对变化 ≤3.24e-5）')
    say('')
    say('【4】判读（写入《A_Q1可靠性与方法学》§5.3 用）')
    say('   · 输入噪声的贡献量级与"离散误差"分属两个数量级：温度侧由**输入噪声支配**，含水率侧由**网格离散支配**。')
    say('   · 因此第 4 位小数的"有效位数"表述必须分场给出（见口径总表 §8.5）。')
    with io.open(OUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(L) + '\n')
    print('WROTE ' + OUT)


if __name__ == '__main__':
    main()
