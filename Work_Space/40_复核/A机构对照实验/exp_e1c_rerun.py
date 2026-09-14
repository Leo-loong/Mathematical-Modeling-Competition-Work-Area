# -*- coding: utf-8 -*-
"""
按《q1_e1_gci.py》的**精确代码路径**重测 E1-c 时间收敛序列，判定登记值是否可复现。
只读调用 q1_core；不覆盖 Q1/code 下的任何日志。
"""
import io
import os
import sys
import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from exp_common import find_root

ROOT = find_root()
sys.path.insert(0, os.path.join(ROOT, '11_建模', '11-3_算法与管线', 'Q1', 'code'))
import q1_core as core          # noqa: E402

ATT1 = os.path.join(ROOT, '10_赛题', 'A题', '附件', '附件1.xlsx')
OUT = os.path.join(HERE, 'exp_e1c_rerun_report.txt')
L = []


def say(s=''):
    print(s, flush=True)
    L.append(str(s))


def load_att1():
    wb = load_workbook(ATT1, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    data = list(ws.iter_rows(values_only=True))[1:]
    return (np.array([float(r[0]) for r in data]),
            np.array([float(r[1]) for r in data]),
            np.array([float(r[2]) for r in data]))


def one_case(N, dr, sub, subT, t, Tinf, Cinf, p, case='collect'):
    """与 q1_e1_gci.one_case 完全同构（cols 取 21 个公共输出点）"""
    cols = np.arange(0, N + 1, max(1, N // 20))
    if case == 'collect':
        res, _, _ = core.run_sim_collect(
            N=N, dr=dr, dt_out=1.0, nsteps=1800, p=p,
            Tenv_fn=lambda s: float(np.interp(s, t, Tinf)),
            Cenv_fn=lambda s: float(np.interp(s, t, Cinf)),
            sub=sub, subT=subT, center='fv', surf='fvm', cols=cols,
            solve_T=False, solve_C=True)
        Cend = res['C_end']
    else:
        res = core.run_sim(
            N=N, dr=dr, dt_out=1.0, nsteps=1800, p=p,
            Tenv_fn=lambda s: float(np.interp(s, t, Tinf)),
            Cenv_fn=lambda s: float(np.interp(s, t, Cinf)),
            sub=sub, subT=subT, center='fv', surf='fvm',
            solve_T=False, solve_C=True)
        Cend = res['C_end']
    return float(Cend[cols][-1]), float(Cend[cols][0]), res['stats']


def main():
    t, Tinf, Cinf = load_att1()
    p = dict(core.DEF)
    say('=' * 92)
    say('E1-c 时间收敛序列 · 按 q1_e1_gci.py 精确路径重测（Δr=0.25 mm，N=80）')
    say('=' * 92)
    say('登记值（《A_Q1求解记录》§6.1(c) 与《A_Q数值口径总表》§8.3 E1-c）：')
    say('   dt_int=1 s    → 1.5104056508')
    say('   dt_int=0.5 s  → 1.5103626682')
    say('   dt_int=0.25 s → 1.5103412843')
    say('   dt_int=0.125 s→ 1.5103338247')
    say('   口径表另写 dt=0.125 s → 1.510434（疑为 1.510334 的笔误）')
    say('')
    say(f'{"调用":<12s} {"sub":>4} {"subT":>5} {"Δt_int/s":>9} {"C(R,1800)":>16} {"C(0,1800)":>14} {"avg_it":>7}')
    for case in ('collect', 'sim'):
        for sub, subT in [(1, 1), (2, 1), (4, 1), (8, 1), (10, 32)]:
            cs, c0, st = one_case(80, 0.25e-3, sub, subT, t, Tinf, Cinf, p, case)
            say(f'{case:<12s} {sub:4d} {subT:5d} {1.0/sub:9.4f} {cs:16.10f} {c0:14.8f} {st["avg_it"]:7.2f}')
    say('')
    say('对照：同一路径在 Δr=0.2 mm（N=100）下的复算（用于验证登记值的真实来源）')
    for sub in (1, 2, 4, 8, 10):
        cs, c0, st = one_case(100, 0.2e-3, sub, 1, t, Tinf, Cinf, p, 'collect')
        say(f'   N=100  sub={sub:2d}  Δt_int={1.0/sub:6.3f} s   C(R,1800)={cs:.10f}')
    with io.open(OUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(L) + '\n')
    print('WROTE ' + OUT)


if __name__ == '__main__':
    main()
