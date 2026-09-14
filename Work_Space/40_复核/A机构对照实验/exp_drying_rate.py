# -*- coding: utf-8 -*-
"""
B2：干燥速率曲线 u(t) = −dC̄/dt 图数据导出（零新增求解，只读 result3.xlsx）
==========================================================================
- 截面平均含水率 C̄(t)：按环形元体体积加权（21 个等距径向点，Δr=1 mm）
- 干燥速率 u(t)：对 C̄(t) 做中心差分（单位 kg/(kg·h)）
- 同时给"分段平均速率"（便于论文读数）：0–24 h / 24–57.53 h

输出：20_交付包/04_图表包/data/fig_q3_drying_rate.csv
"""
import io
import os
import csv
import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
import sys
sys.path.insert(0, HERE)
from exp_common import find_root

ROOT = find_root()
SRC = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'results', 'result3.xlsx')
OUT = os.path.join(ROOT, '20_交付包', '04_图表包', 'data', 'fig_q3_drying_rate.csv')
R0 = 0.02          # m
NR = 20            # 0.0 … 2.0 cm，共 21 列
DR = R0 / NR


def main():
    wb = load_workbook(SRC, read_only=True, data_only=True)
    ws = wb['Sheet1']
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = rows[0]
    body = [r for r in rows[1:] if r[0] is not None]
    t = np.array([float(r[0]) for r in body])
    C = np.array([[float(v) for v in r[1:1 + NR + 1]] for r in body])   # (nt, 21)
    n = len(t)

    # 环形元体体积权重（以 πR0² 归一）
    r = np.arange(NR + 1) * DR
    rf = np.concatenate([[0.0], 0.5 * (r[:-1] + r[1:]), [R0]])          # 面半径，共 NR+2
    V = np.array([rf[i + 1] ** 2 - rf[i] ** 2 for i in range(NR + 1)]) / R0 ** 2
    Cbar = C @ V / V.sum()

    # 干燥速率（中心差分，端点用单侧差分）
    u = np.zeros(n)
    u[1:-1] = -(Cbar[2:] - Cbar[:-2]) / (t[2:] - t[:-2])
    u[0] = -(Cbar[1] - Cbar[0]) / (t[1] - t[0])
    u[-1] = -(Cbar[-1] - Cbar[-2]) / (t[-1] - t[-2])
    u_h = u * 3600.0                                                    # kg/(kg·h)

    print('rows=%d  t=%.0f…%.0f s' % (n, t[0], t[-1]))
    print('C̄(0)=%.6f  C̄(end)=%.6f  t_end=%.4f h' % (Cbar[0], Cbar[-1], t[-1] / 3600))
    print('u(0)=%.4f  u(24h)=%.4f  u(end)=%.4f  kg/(kg·h)'
          % (u_h[0], u_h[np.argmin(abs(t - 86400))], u_h[-1]))
    m24 = t <= 86400
    print('0–24 h 平均速率 = %.4f kg/(kg·h)（脱除 %.4f）'
          % ((Cbar[0] - Cbar[m24][-1]) / (86400 / 3600), Cbar[0] - Cbar[m24][-1]))
    m2 = t > 86400
    print('24 h–结束 平均速率 = %.4f kg/(kg·h)（脱除 %.4f）'
          % ((Cbar[m24][-1] - Cbar[-1]) / ((t[-1] - 86400) / 3600), Cbar[m24][-1] - Cbar[-1]))

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with io.open(OUT, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(['t_s', 't_h', 'Cbar', 'u_kg_per_kg_s', 'u_kg_per_kg_h'])
        for i in range(n):
            w.writerow(['%.0f' % t[i], '%.4f' % (t[i] / 3600), '%.8f' % Cbar[i],
                        '%.6e' % u[i], '%.6f' % u_h[i]])
    print('WROTE', OUT)


if __name__ == '__main__':
    main()
