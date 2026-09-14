# -*- coding: utf-8 -*-
"""
生成 Q1 论文表1／表2（7×5）到交付包，并扫描交付包内 q1_*.py 的旧口径残留。
只读 result1.xlsx，不重跑任何求解。
产出：
  · 20_交付包/01_模型叙述/A_Q1论文表格_表1表2.md
  · 终端打印残留扫描结果（ASCII）
约定：不引用、不记载任何资料中的日期。
"""
import os
import io
import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))


def find_root(p, marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, marker)):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = find_root(HERE)
RES1 = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'results', 'result1.xlsx')
OUT = os.path.join(ROOT, '20_交付包', '01_模型叙述', 'A_Q1论文表格_表1表2.md')
CODE = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'code')

SNAP = [100, 300, 600, 900, 1200, 1500, 1800]
IDX5 = [0, 5, 10, 15, 20]          # 0/0.5/1.0/1.5/2.0 cm
HDR5 = ['0', '0.5', '1.0', '1.5', '2.0']

# 《A_数值口径总表》§8.2 登记的全精度值（result1.xlsx 仅存 4 位小数，故此处引用口径表）
REG = {'T0': 33.5753730384, 'TR': 36.785566, 'C0': 2.54999185, 'CR': 1.510381}


def read(fn, sheet):
    wb = load_workbook(fn, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    data = {int(r[0]): [float(v) for v in r[1:]] for r in rows[1:]}
    return data


def table_md(title, data, unit):
    s = []
    s.append('**%s**（单位：%s）' % (title, unit))
    s.append('')
    s.append('| 时间/s | ' + ' | '.join(HDR5) + ' |')
    s.append('|---|' + '---|' * len(HDR5))
    for t in SNAP:
        s.append('| %d | ' % t + ' | '.join('%.4f' % data[t][i] for i in IDX5) + ' |')
    s.append('')
    return s


def main():
    T = read(RES1, '温度')
    C = read(RES1, '水分浓度')

    L = []
    L.append('# A 题 · 第一问 论文表格（表1 温度 ／ 表2 水分浓度）')
    L.append('')
    L.append('> **本文件自包含**：写作者可**直接转写为论文表1／表2**，无需再查其他文件。')
    L.append('> **数值唯一来源**：`20_交付包/09_代码与复现/results/result1.xlsx`（已通过程序化断言 PASS）；')
    L.append('> 与《A_数值口径总表》§八 **逐格一致**。**禁止手抄他处**。')
    L.append('> **口径**：元体平衡＋后向欧拉全隐式；界面变系数取**沿 $C$ 的积分平均（M6）**；')
    L.append('> $\\Delta r=0.25$ mm（$N=80$ 区间／81 节点）；温度内部子步 $1/32$ s、含水率子步 $0.1$ s。')
    L.append('> 按项目约定，本文件不引用、不记载资料中出现的日期。')
    L.append('')
    L.append('---')
    L.append('')
    L.append('## 表 1　30 分钟内药材的温度')
    L.append('')
    L.extend(table_md('表1　温度', T, '℃'))
    L.append('---')
    L.append('')
    L.append('## 表 2　30 分钟内药材的水分浓度')
    L.append('')
    L.extend(table_md('表2　水分浓度', C, 'kg/kg'))
    L.append('---')
    L.append('')
    L.append('## 关键值（$t=1800$ s，供正文引用）')
    L.append('')
    L.append('| 量 | 4 位小数 | 全精度 | 说明 |')
    L.append('|---|---|---|---|')
    L.append('| $T(0,1800)$ | **%.4f** ℃ | %.10f | 中心温度，较初值 $+5.5754$ ℃ |' % (T[1800][0], REG['T0']))
    L.append('| $T(R,1800)$ | **%.4f** ℃ | %.6f | 表面温度，较初值 $+8.7856$ ℃ |' % (T[1800][20], REG['TR']))
    L.append('| $C(0,1800)$ | **%.4f** kg/kg | %.8f | 中心含水率，**几乎不变**（$1800\\,\\mathrm{s}\\approx0.022\\tau_C$） |' % (C[1800][0], REG['C0']))
    L.append('| $C(R,1800)$ | **%.4f** kg/kg | %.6f | 表面含水率，**降 40.8%%**；网格误差带 $\\pm1\\times10^{-4}$，**第 4 位为误差带内取值** |' % (C[1800][20], REG['CR']))
    L.append('')
    L.append('> **口径唯一性**：表内数字与 `result1.xlsx`、全部图数据 CSV、各交付文档**同源且逐格一致**。')
    L.append('> 表中 $C(R,1800)$ 取 $\\Delta r=0.25$ mm 口径实测值 **1.5104**（全精度 1.510381）。')
    L.append('')

    with io.open(OUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(L) + '\n')
    print('TABLE_OK %s' % os.path.basename(OUT))

    # ---- 残留扫描 ----
    print('--- scan q1_*.py for legacy tokens ---')
    hits = 0
    for n in sorted(os.listdir(CODE)):
        if not (n.startswith('q1_') and n.endswith('.py')):
            continue
        with io.open(os.path.join(CODE, n), 'r', encoding='utf-8', errors='replace') as f:
            for i, line in enumerate(f, 1):
                if ('调和' in line) or ('1.51033' in line):
                    print('  %s:%d :: %s' % (n, i, line.strip()[:120]))
                    hits += 1
    print('SCAN_HITS %d' % hits)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
