# -*- coding: utf-8 -*-
"""
Q1 唯一基线审计（只读）
=======================
1) 对比 两份 code 目录（工作区 11-3/Q1/code  vs  交付包 09/code）的 q1_* 脚本是否逐字节一致
2) 对比 交付包 04_图表包/data/*.csv 与 30_图表/03_图数据准备/*.csv 是否一致
3) 读 result1.xlsx，输出 t=1800 关键值与表1/表2 抽样
4) 扫 q1_* 脚本内是否残留旧登记值 1.51033 / 1.510332
只输出报告到 exp_q1_baseline_audit_report.txt（UTF-8），终端仅 ASCII 状态。
"""
import os
import io
import hashlib
import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))


def find_root(p, marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, marker)):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = find_root(HERE)
WS_CODE = os.path.join(ROOT, '11_建模', '11-3_算法与管线', 'Q1', 'code')
DP_CODE = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'code')
DP_DATA = os.path.join(ROOT, '20_交付包', '04_图表包', 'data')
WS_DATA = os.path.join(ROOT, '30_图表', '03_图数据准备')
RES1 = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'results', 'result1.xlsx')
REP = os.path.join(HERE, 'exp_q1_baseline_audit_report.txt')

L = []


def say(s):
    L.append(str(s))


def md5(fn):
    h = hashlib.md5()
    with open(fn, 'rb') as f:
        h.update(f.read())
    return h.hexdigest()


def main():
    say('=' * 84)
    say('Q1 唯一基线审计（只读）')
    say('ROOT = %s' % ROOT)
    say('=' * 84)

    # ---- 1. 代码副本一致性 ----
    say('')
    say('--- 1. 两份 code 目录的 q1_* 脚本 hash 对比 ---')
    names = sorted(n for n in os.listdir(WS_CODE) if n.startswith('q1_') and n.endswith('.py'))
    same = diff = miss = 0
    for n in names:
        pw = os.path.join(WS_CODE, n)
        pd = os.path.join(DP_CODE, n)
        if not os.path.exists(pd):
            say('  [仅工作区] %s' % n)
            miss += 1
            continue
        hw, hd = md5(pw), md5(pd)
        if hw == hd:
            same += 1
        else:
            say('  [不一致]   %s' % n)
            diff += 1
    say('  合计 %d 个：一致 %d / 不一致 %d / 仅工作区 %d' % (len(names), same, diff, miss))

    # ---- 2. 图数据 CSV 双份一致性 ----
    say('')
    say('--- 2. 图数据 CSV 双份对比（交付包/04_图表包/data vs 30_图表/03_图数据准备）---')
    csvs = sorted(n for n in os.listdir(WS_DATA) if n.endswith('.csv'))
    for n in csvs:
        pw = os.path.join(WS_DATA, n)
        pd = os.path.join(DP_DATA, n)
        if not os.path.exists(pd):
            say('  [缺交付副本] %s' % n)
            continue
        hw, hd = md5(pw), md5(pd)
        say('  %-28s %s' % (n, '一致' if hw == hd else '★不一致'))
    only_dp = sorted(set(os.listdir(DP_DATA)) - set(csvs))
    if only_dp:
        say('  [仅交付包] %s' % ', '.join(only_dp))

    # ---- 3. result1.xlsx 关键值 ----
    say('')
    say('--- 3. result1.xlsx 现状 ---')
    wb = load_workbook(RES1, data_only=True, read_only=True)
    say('  sheets = %s' % wb.sheetnames)
    d = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [float(x) for x in rows[0][1:]]
        t = [int(r[0]) for r in rows[1:]]
        A = np.array([[float(v) for v in r[1:]] for r in rows[1:]])
        d[sn] = (hdr, t, A)
        say('  [%s] rows=%d cols=%d  t=[%d..%d]' % (sn, len(rows), ws.max_column, t[0], t[-1]))
    hT, tT, AT = d['温度']
    hC, tC, AC = d['水分浓度']
    i5 = [0, 5, 10, 15, 20]
    say('  表1（温度）末行 t=1800：' + '  '.join('%.4f' % AT[-1, j] for j in i5))
    say('  表2（水分浓度）末行 t=1800：' + '  '.join('%.4f' % AC[-1, j] for j in i5))
    say('  距离列 = %s' % ('0..2.0 (21 列)' if len(hT) == 21 else str(len(hT))))

    # ---- 4. 脚本内旧登记值残留 ----
    say('')
    say('--- 4. q1_* 脚本内旧登记值残留扫描（1.51033 / 1.510332）---')
    pats = ['1.51033', '1.510332']
    for base in (WS_CODE, DP_CODE):
        for n in sorted(os.listdir(base)):
            if not (n.startswith('q1_') and n.endswith('.py')):
                continue
            p = os.path.join(base, n)
            with io.open(p, 'r', encoding='utf-8', errors='replace') as f:
                for i, line in enumerate(f, 1):
                    for pt in pats:
                        if pt in line:
                            say('  %s :: %s:%d :: %s' % (
                                '工作区' if base == WS_CODE else '交付包', n, i, line.strip()))

    with io.open(REP, 'w', encoding='utf-8') as f:
        f.write('\n'.join(L) + '\n')
    print('AUDIT DONE -> exp_q1_baseline_audit_report.txt')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
