# -*- coding: utf-8 -*-
"""
A 题 Q1（预热平衡阶段 0–1800 s）主力求解器 v4
==============================================
方法：空间＝元体平衡（有限体积）＋ 三对角追赶法；
      时间＝**温度侧时间方向精确推进**（`q1_exact.run_T_expm`，h=1 s、无时间离散误差）
            ＋ 含水率侧全隐式（后向欧拉，子步 0.1 s）
物性：附录2（常物性 ＋ D(C) 单因子非线性）
中心：两条**独立推导**的装配路径（元体平衡几何式 / 洛必达极限式），做解级比对
表面：元体平衡半格（Robin）

相对 v3 的变更（据《A_Q1求解记录》§九 待办）
--------------------------------------------
1. **单遍运行**：同一次求解同时产出整表输出与论文快照，
   消除"表1/表2 与 result1.xlsx 来自两遍计算"的异源风险；
2. **中心双路真装配**：v3 中两条分支写成了同一表达式的两遍（断言自证），
   本版改为两条独立推导（几何式 vs 极限式）并做**解级**比对；
3. 求解核抽离至 `q1_core.py`（配置化），供 E1–E5 复用；
4. 清理死代码 `write_and_assert`；日志补运行环境与耗时分解。

运行（**须显式 `--go`**，见《工作约束》§12.1 运行闸门）：
  python q1_solver.py --go
不加 --go 时**只打印 [SCALE] 规模声明并退出**，不执行求解。

约定：本文件不引用、不记载任何资料中的日期。
"""
import os
import io
import sys
import time
import platform
import openpyxl
import numpy as np
from openpyxl import load_workbook, Workbook

import q1_core as core
import q1_exact as ex
from q1_core import DEF

HERE = os.path.dirname(os.path.abspath(__file__))
# 自适应定位工作区根：向上查找含 10_赛题 的目录
# （兼容两种目录深度：工作区 11_建模/11-3_算法与管线/Q1/code/ 与交付包 09 的 code/）
def _find_root(p, _marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, _marker)):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = _find_root(HERE)
ATT1 = os.path.join(ROOT, '10_赛题', 'A题', '附件', '附件1.xlsx')
LOG = os.path.join(HERE, 'q1_run_log.txt')
RESDIR = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'results')
RES1 = os.path.join(RESDIR, 'result1.xlsx')

DR = 0.00025        # 空间步长（m）
DT = 1.0            # 输出步长（s）
SUB = 10            # C 方程内部子步数（内部步长 = 0.1 s）
SUBT = 32           # **后向欧拉对照路径**下 T 方程的内部子步数（内部步长 = 1/32 s）
#   ↑ ⚠ **交付路径不使用本参数**：温度侧已改为「时间方向精确推进」（hstep = DT = 1 s，无内部子步）；
#     本参数仅供"后向欧拉对照口径"（E1-b 时间收敛实验等）使用。
#     该旧口径下 dtT = 1 s 存在 ~1.2e-3 ℃ 一阶误差、细化到 1/32 s 仍残留 ~4e-5 ℃，
#     这正是改用精确推进的原因（步长无关性：h=1 s 与 h=60 s 差 3.1e-11 ℃）。
OMEGA = 0.7         # Picard 欠松弛
TOL = 1e-9          # 相对残差判据
MAXIT = 30
NSTEP = 1800
SNAP = (100, 300, 600, 900, 1200, 1500, 1800)
LINES = []


def say(s):
    LINES.append(str(s))
    print(s)


def load_att1():
    wb = load_workbook(ATT1, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    data = list(ws.iter_rows(values_only=True))[1:]
    return (np.array([float(r[0]) for r in data]),
            np.array([float(r[1]) for r in data]),
            np.array([float(r[2]) for r in data]))


def assert_result1():
    """程序化格式断言：工作表名 / 维度 / 表头 / 时间列 / 4 位小数"""
    wb = load_workbook(RES1, read_only=True)
    msgs = []
    ok = True
    if wb.sheetnames != ['温度', '水分浓度']:
        ok = False
        msgs.append('FAIL sheetnames=%s' % wb.sheetnames)
    else:
        msgs.append('OK sheetnames=%s' % wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        msgs.append('%s rows=%d cols=%d' % (sn, len(rows), ws.max_column))
        if len(rows) != 1801 or ws.max_column != 22:
            ok = False
            msgs.append('FAIL dims (expect 1801x22)')
        if rows[0][0] != '时间\\到药材中心的距离':
            ok = False
            msgs.append('FAIL header A1')
        if [round(float(x), 4) for x in rows[0][1:]] != [round(0.1 * k, 4) for k in range(21)]:
            ok = False
            msgs.append('FAIL header cols')
        else:
            msgs.append('OK header cols 0..2.0 (21)')
        if [r[0] for r in rows[1:]] != list(range(1, 1801)):
            ok = False
            msgs.append('FAIL time col')
        else:
            msgs.append('OK time col 1..1800')
        bad = sum(1 for r in rows[1:] for v in r[1:]
                  if v is None or abs(round(float(v), 4) - float(v)) > 1e-12)
        if bad:
            ok = False
            msgs.append('FAIL not 4-decimal: %d' % bad)
        else:
            msgs.append('OK all cells 4-decimal')
    return ok, msgs


def main():
    t_all = time.time()
    t, Tinf, Cinf = load_att1()
    p = dict(DEF)
    N = int(round(p['R0'] / DR))
    cols = np.arange(0, N + 1, 4)                 # 0.1 cm 步长 → 21 列
    Tenv = lambda s: float(np.interp(s, t, Tinf))
    Cenv = lambda s: float(np.interp(s, t, Cinf))

    # ---------------------------------------------------------------- W0
    say('=== W0 附件核验与平台统计 ===')
    say('points=%d  t=[%g,%g]  dt=%g' % (len(t), t[0], t[-1], t[1] - t[0]))
    say('Tinf first/last=%.4f/%.4f  Cinf first/last=%.5f/%.5f'
        % (Tinf[0], Tinf[-1], Cinf[0], Cinf[-1]))
    seg = (t >= 10000) & (t <= 14400)
    say('plateau[10000,14400]: Tinf mean=%.4f range=%.4f ; Cinf mean=%.6f range=%.6f'
        % (Tinf[seg].mean(), Tinf[seg].max() - Tinf[seg].min(),
           Cinf[seg].mean(), Cinf[seg].max() - Cinf[seg].min()))
    say('env: python=%s  numpy=%s  openpyxl=%s  %s %s'
        % (platform.python_version(), np.__version__, openpyxl.__version__,
           platform.system(), platform.release()))

    # ------------------------------------------------- W2 中心双路（真装配）
    say('')
    say('=== W2 中心节点双路独立装配与解级比对 ===')
    a1, b1, c1, m1 = core.assemble_T(N, DR, DT, p, 'fv', 'fvm')
    a2, b2, c2, m2 = core.assemble_T(N, DR, DT, p, 'lh', 'fvm')
    row_ok = bool(np.allclose([b1[0], c1[0]], [b2[0], c2[0]], rtol=0.0, atol=1e-12))
    say('归一化中心行  fv: b0=%.12e c0=%.12e' % (b1[0], c1[0]))
    say('归一化中心行  lh: b0=%.12e c0=%.12e' % (b2[0], c2[0]))
    say('中心行逐系数一致 = %s' % row_ok)

    rT_fv = core.run_sim(N, DR, DT, NSTEP, p, Tenv, Cenv, sub=SUB, subT=SUBT,
                         center='fv', surf='fvm', solve_C=False)
    rT_lh = core.run_sim(N, DR, DT, NSTEP, p, Tenv, Cenv, sub=SUB, subT=SUBT,
                         center='lh', surf='fvm', solve_C=False)
    dmax = float(np.max(np.abs(rT_fv['T_end'] - rT_lh['T_end'])))
    bitwise = bool(np.array_equal(rT_fv['T_end'], rT_lh['T_end']))
    say('解级比对：两条中心装配的 T_end 最大逐点差 = %.3e K（逐位相同=%s）' % (dmax, bitwise))

    # ------------------------------------------------------------ 主力求解
    say('')
    say('=== 主力求解（T：时间精确推进；C：隐式＋Picard；整表与快照同源） ===')
    t0 = time.time()
    rT = ex.run_T_expm(N, DR, p, Tenv, NSTEP * DT, hstep=DT, center='fv', surf='fvm',
                       t_snaps=set(SNAP), cols=cols, dt_out=DT)
    tT = time.time() - t0
    tabT = rT['table']
    resC, _, tabC = core.run_sim_collect(
        N=N, dr=DR, dt_out=DT, nsteps=NSTEP, p=p,
        Tenv_fn=Tenv, Cenv_fn=Cenv, sub=SUB, subT=SUBT,
        center='fv', surf='fvm', tol=TOL, omega=OMEGA, maxit=MAXIT,
        cols=cols, snap_at=set(SNAP), solve_T=False, solve_C=True)
    st = resC['stats']
    say('N=%d  dr=%.3g m  节点数=%d  输出列=%d' % (N, DR, N + 1, len(cols)))
    say('T 精确推进 %.2f s（%d 步／传播子 %d 个，h=%.3g s）; C 隐式 %.1f s（%d 子步，dt_int=%.3g s）'
        % (tT, rT['nstep'], rT['nprop'], rT['hstep'], resC['C_time'], st['nsub'], resC['dt_int']))
    say('C Picard: avg_iters=%.2f  max_iters=%d  max_res=%.2e'
        % (st['avg_it'], st['max_it'], st['max_res']))
    Tend, Cend = rT['T_end'], resC['C_end']
    say('T center/surface=%.6f/%.6f   C center/surface=%.6f/%.6f'
        % (Tend[0], Tend[-1], Cend[0], Cend[-1]))

    # ------------------------------------------------------------ W5 输出
    say('')
    say('=== W5 输出与程序化断言 ===')
    os.makedirs(RESDIR, exist_ok=True)
    wb = Workbook()
    wsT = wb.active
    wsT.title = '温度'
    wsC = wb.create_sheet('水分浓度')
    for ws in (wsT, wsC):
        ws.cell(row=1, column=1, value='时间\\到药材中心的距离')
        for j in range(21):
            ws.cell(row=1, column=2 + j, value=round(0.1 * j, 4))
    for n in range(1, NSTEP + 1):
        wsT.cell(row=1 + n, column=1, value=n)
        wsC.cell(row=1 + n, column=1, value=n)
        for j in range(21):
            wsT.cell(row=1 + n, column=2 + j, value=round(float(tabT[n - 1, j]), 4))
            wsC.cell(row=1 + n, column=2 + j, value=round(float(tabC[n - 1, j]), 4))
    wb.save(RES1)
    say('result1.xlsx written')
    ok, msgs = assert_result1()
    for m in msgs:
        say('  ' + m)
    say('ASSERTION %s' % ('PASS' if ok else 'FAIL'))

    # ------------------------------------------------------- 论文表格
    say('')
    say('=== 表1 温度（℃）／表2 水分浓度（kg/kg），列＝到中心距离 0–2.0 cm ===')
    idx5 = [int(round(x / DR)) for x in (0.0, 0.005, 0.01, 0.015, 0.02)]
    say('Table 1 (T degC)')
    for ts in SNAP:
        say('%5d  ' % ts + '  '.join('%.4f' % rT['T_snap'][float(ts)][i] for i in idx5))
    say('Table 2 (C kg/kg)')
    for ts in SNAP:
        say('%5d  ' % ts + '  '.join('%.4f' % resC['C_snap'][ts][i] for i in idx5))

    say('')
    say('total elapsed=%.1fs' % (time.time() - t_all))
    with io.open(LOG, 'w', encoding='utf-8') as f:
        f.write('\n'.join(LINES) + '\n')
    print('DONE assert=%s' % ok)
    return 0 if ok else 1


# ---------------------------------------------------------------------------
# 运行闸门（《工作约束》§12.1）：默认只"声明规模"并退出；须显式 `--go`
# （或环境变量 ALLOW_RUN=1）且规模不超批准值，才真正执行。
# 首行输出 [SCALE] 声明行，便于事后审计。
# ---------------------------------------------------------------------------
def _gate(name, scale, approved=None, argv=None, env=None):
    args = sys.argv if argv is None else argv
    envv = os.environ if env is None else env
    print('[SCALE] %s | ' % name + ' | '.join('%s=%s' % (k, v) for k, v in scale.items()),
          flush=True)
    if not (('--go' in args) or (envv.get('ALLOW_RUN', '') == '1')):
        print('[GATE] 未传入 --go（或 ALLOW_RUN=1）⟹ 仅声明规模，不执行。', flush=True)
        print('[GATE] 真正执行：python %s --go' % os.path.basename(__file__), flush=True)
        return False
    if approved:
        bad = ['%s %s>%s' % (k, scale[k], approved[k])
               for k in approved if k in scale and scale[k] > approved[k]]
        if bad:
            print('[GATE] 规模超批准值（%s）⟹ 拒绝执行。' % '；'.join(bad), flush=True)
            return False
    print('[GATE] 规模在校验范围内 ⟹ 执行。', flush=True)
    return True


if __name__ == '__main__':
    if not _gate('Q1 主力求解 v4',
                 dict(网格='Δr=0.25 mm（81 节点）', 输出='0–1800 s @1 s（1800 行）',
                      内部子步='T 无（时间精确推进，h=1 s）；C 0.1 s', 预计耗时='约 55 s')):
        sys.exit(0)
    sys.exit(main())
