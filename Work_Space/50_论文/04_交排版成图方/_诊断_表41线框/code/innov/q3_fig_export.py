# -*- coding: utf-8 -*-
"""
Q3 图数据导出（F-33 / F-34）—— 遗留项 L-1
================================================================================
【性质】只读导出：不修改任何求解数值与交付物，仅生成绘图用 CSV。
【产出】写入工作区真源 30_图表/03_图数据准备/（再同步至交付包 data/）
  · fig_q3_history.csv    列：t_h, C_center, C_surface, C_mean_env      （F-33）
  · fig_q3_conv.csv       列：kind, dr_mm, nsub, C0, CR, rel_change      （F-34 左/中）
  · fig_q3_conv_pos.csv   列：pair, loc_cm, rel_change                   （F-34 右）

【数据源】
  · result3.xlsx（中心/表面含水率时程，每 60 s）           —— F-33
  · 附件1.xlsx（环境温湿度时序 ＋ 恒温段常值延拓）          —— F-33 环境列
  · logs/q3_w2_log.txt   （**现行口径**复跑后的时间/空间收敛）  —— F-34 左/中
  · logs/q3_grid_log.txt （**现行口径**复跑后的网格裁决与各位置变化）—— F-34 右

【用法】
  python q3_fig_export.py                  # dry-run（只打印摘要）
  python q3_fig_export.py --go             # 写盘
  python q3_fig_export.py --only f34 --go  # 仅导出 F-34
"""
import sys
import os
import re
import csv

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

for _v in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS'):
    os.environ.setdefault(_v, '1')

import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
Q3ROOT = os.path.abspath(os.path.join(HERE, '..'))
WS = os.path.abspath(os.path.join(HERE, '..', '..', '..', '..'))
sys.path.insert(0, Q3ROOT)

RES3 = os.path.join(WS, '20_交付包', '09_代码与复现', 'results', 'result3.xlsx')
ATT1 = os.path.join(WS, '10_赛题', 'A题', '附件', '附件1.xlsx')
Q3LOG = os.path.join(WS, '11_建模', '11-3_算法与管线', 'Q3', 'logs')
W2LOG = os.path.join(Q3LOG, 'q3_w2_log.txt')
GRLOG = os.path.join(Q3LOG, 'q3_grid_log.txt')
OUTD = os.path.join(WS, '30_图表', '03_图数据准备')

GO = '--go' in sys.argv
ONLY = None
for _i, _a in enumerate(sys.argv):
    if _a == '--only' and _i + 1 < len(sys.argv):
        ONLY = sys.argv[_i + 1].lower()

print('[SCALE] 只读导出：result3.xlsx / 附件1.xlsx / q3_w2_log.txt / q3_grid_log.txt')
print('[SCALE] 写盘目标：%s' % OUTD)
print('[SCALE] 模式：%s%s' % ('GO（写盘）' if GO else 'dry-run（不写盘）',
                              '' if ONLY is None else '  仅 ' + ONLY))
print('=' * 78)


# ------------------------------------------------------------------ #
def dump(path, header, rows, fmt=None):
    """写 CSV 或 dry-run 打印摘要。"""
    if not GO:
        print('  [dry-run] %s  ← %d 行' % (os.path.basename(path), len(rows)))
        for r in rows[:3]:
            print('            %s' % (r,))
        print('            ...')
        if len(rows) > 3:
            print('            %s' % (rows[-1],))
        return
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in rows:
            w.writerow(r)
    print('  ✔ 已写 %s（%d 行）' % (path, len(rows)))


# ============================== F-33 ============================== #
def build_f33():
    import q3_core as q3

    print('\n【F-33】中心/表面含水率长时程 ＋ 环境列')
    wb = openpyxl.load_workbook(RES3, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    tv = np.array([float(r[0]) for r in rows[1:]])
    C0 = np.array([float(r[1]) for r in rows[1:]])       # 0.0 cm（中心）
    CR = np.array([float(r[21]) for r in rows[1:]])      # 2.0 cm（表面）
    print('  源 result3.xlsx：%d 行；末行 t=%.0f s (%.4f h)  C_center=%.4f  C_surface=%.4f'
          % (len(tv), tv[-1], tv[-1] / 3600.0, C0[-1], CR[-1]))

    # 环境浓度时程（与求解核**同口径**：附件1 时序 ＋ 恒温段常值延拓 50.00 ℃ / 0.04999）
    wb1 = openpyxl.load_workbook(ATT1, data_only=True)
    ws1 = wb1.active
    r1 = [r for r in ws1.iter_rows(values_only=True)][1:]
    d1 = [r for r in r1 if r[0] is not None]
    ts, Tv, Cv = q3.env_tables_from_att1(
        np.array([float(r[0]) for r in d1]),
        np.array([float(r[1]) for r in d1]),
        np.array([float(r[2]) for r in d1]))
    wb1.close()
    Cenv = np.array([float(q3._interp1_nb(t, ts, Cv)) for t in tv])
    print('  环境列：附件1 时序（至 %.0f s）＋ 恒温段常值延拓；末值 C_env=%.5f' % (ts[-1], Cenv[-1]))

    out = [(round(t / 3600.0, 6), round(c0, 6), round(cR, 6), round(ce, 6))
           for t, c0, cR, ce in zip(tv, C0, CR, Cenv)]
    dump(os.path.join(OUTD, 'fig_q3_history.csv'),
         ['t_h', 'C_center', 'C_surface', 'C_mean_env'], out)

    # 自检：末行必须为达标临界值
    assert abs(C0[-1] - 0.15) < 5e-4, '末行 C_center 应 ≈ 0.1500，实际 %.6f' % C0[-1]
    print('  自检：末行 C_center = %.4f（达标临界 0.1500）✔' % C0[-1])


# ============================== F-34 ============================== #
RE_SPACE = re.compile(r'^\s+(\d+\.\d+)\s+(\d+)\s+([\d.eE+-]+)\s+([\d.eE+-]+)\s')
RE_TIME = re.compile(r'^\s+(\d+\.\d+)\s+([\d.eE+-]+)\s+([\d.eE+-]+)\s')
RE_RELS = re.compile(r'(\d+\.\d+) vs (\d+\.\d+) mm[：:]\s*\|ΔC\(0\)\|/C\(0\)=([\d.eE+-]+)\s+\|ΔC\(R\)\|/C\(R\)=([\d.eE+-]+)')
RE_RELT = re.compile(r'h=1/(\d+)\s+s[：:]\s*\|ΔC\(0\)\|/C\(0\)=([\d.eE+-]+)\s+\|ΔC\(R\)\|/C\(R\)=([\d.eE+-]+)')
RE_GRIDHEAD = re.compile(r'Δr=([\d.]+)\s+mm\s+N=(\d+)')
RE_GRIDPOS = re.compile(r'^\s+(\d+\.\d+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s*$')
RE_GRIDREL = re.compile(r'(\d+\.\d+) vs (\d+\.\d+)\s+(-?[\d.]+)\s+(-?[\d.]+)\s+(-?[\d.]+)\s+(-?[\d.]+)\s+(-?[\d.]+)')


def parse_w2(path):
    txt = open(path, encoding='utf-8').read().splitlines()
    sec = None
    space, tim = [], []
    rels, relt = {}, {}
    for ln in txt:
        if '① 空间收敛' in ln:
            sec = 's'
            continue
        if '② 时间收敛' in ln:
            sec = 't'
            continue
        if '相邻网格相对变化（以更细者为基准）' in ln:
            sec = 'rs'
            continue
        if '以内部步 1/32 s 为基准的相对变化' in ln:
            sec = 'rt'
            continue
        if sec == 's':
            m = RE_SPACE.match(ln)
            if m:
                space.append((float(m.group(1)), int(m.group(2)),
                              float(m.group(3)), float(m.group(4))))
        elif sec == 't':
            m = RE_TIME.match(ln)
            if m:
                tim.append((float(m.group(1)), float(m.group(2)), float(m.group(3))))
        elif sec == 'rs':
            m = RE_RELS.search(ln)
            if m:
                rels[float(m.group(1))] = float(m.group(3))   # 更细档 → C(0) 相对变化
        elif sec == 'rt':
            m = RE_RELT.search(ln)
            if m:
                relt[int(m.group(1))] = float(m.group(2))     # 1/N → C(0) 相对变化
    return space, tim, rels, relt


def parse_grid(path):
    txt = open(path, encoding='utf-8').read().splitlines()
    sec = None
    c0map = {}
    pos = {}
    rel = []
    for ln in txt:
        m = RE_GRIDHEAD.search(ln)
        if m:
            c0map[float(m.group(1))] = None
            continue
        if '各位置的值' in ln:
            sec = 'p'
            continue
        if '相邻网格的相对变化' in ln:
            sec = 'r'
            continue
        if sec == 'p':
            m = RE_GRIDPOS.match(ln)
            if m:
                pos[float(m.group(1))] = [float(m.group(i)) for i in range(2, 7)]
        elif sec == 'r':
            m = RE_GRIDREL.search(ln)
            if m:
                rel.append((float(m.group(1)), float(m.group(2)),
                            [float(m.group(i)) for i in range(3, 8)]))
    return c0map, pos, rel


def build_f34():
    print('\n【F-34】长时程收敛性与网格裁决')
    for p in (W2LOG, GRLOG):
        if not os.path.exists(p):
            print('  ✘ 缺少日志：%s（请先完成 L-6 复跑）' % p)
            return

    space, tim, rels, relt = parse_w2(W2LOG)
    print('  解析 q3_w2_log.txt：空间 %d 档、时间 %d 档、空间相对变化 %d 条、时间相对变化 %d 条'
          % (len(space), len(tim), len(rels), len(relt)))

    rows = []
    for dr, N, c0, cR in space:
        rows.append(['space', '%.5f' % dr, 32, '%.6f' % c0, '%.6f' % cR,
                     ('%.6e' % rels[dr]) if dr in rels else ''])
    for h_in, c0, cR in tim:
        nsub = int(round(1.0 / h_in))
        rows.append(['time', '0.25000', nsub, '%.6f' % c0, '%.6f' % cR,
                     ('%.6e' % relt[nsub]) if nsub in relt else ''])
    dump(os.path.join(OUTD, 'fig_q3_conv.csv'),
         ['kind', 'dr_mm', 'nsub', 'C0', 'CR', 'rel_change'], rows)

    # -------- 各位置（F-34 右图）-------- #
    c0map, pos, rel = parse_grid(GRLOG)
    print('  解析 q3_grid_log.txt：位置表 %d 档、网格对 %d 组' % (len(pos), len(rel)))
    LOCs = [0.0, 0.5, 1.0, 1.5, 2.0]
    prows = []
    for fine, coarse, vals in rel:
        for j, loc in enumerate(LOCs):
            prows.append(['%.5f vs %.5f' % (fine, coarse), loc, '%.3f' % vals[j]])
    dump(os.path.join(OUTD, 'fig_q3_conv_pos.csv'),
         ['pair', 'loc_cm', 'rel_change'], prows)

    # 自检：与《A_数值口径总表》§10.3 对照
    print('\n  --- 与口径表 §10.3 对照（人工核对）---')
    for r in rows:
        print('    %-6s dr=%-8s nsub=%-4s C0=%-10s CR=%-10s rel=%s' % tuple(r))


# ------------------------------------------------------------------ #
if ONLY in (None, 'f33'):
    build_f33()
if ONLY in (None, 'f34'):
    build_f34()

print('\n' + '=' * 78)
if GO:
    print('完成。请同步至 20_交付包/04_图表包/data/ 后再更新规格卡与交付清单。')
else:
    print('dry-run 结束（未写盘）。加 --go 才会写盘。')
