# -*- coding: utf-8 -*-
"""
Q3 · 表面加密网格的两项收尾验证（只验证，不产交付物）
================================================================================
① **时间收敛**：在细网格（表面加密档 M=26）上比较内部步 1/32 s 与 1/64 s（比较时刻 24 h）。
   —— 细网格含极小单元，须确认时间离散仍未成为误差主导。
② **Q2 重叠段回归**：细网格在 0–3 h 的 C(0) 与 `result2.xlsx` 比对（逐时刻快照）。
   —— 早期 δ≫Δr，表面加密不应改变早期解；须与"逐位一致"结论相容。

用法：
  python q3_verify_grid.py                         # 规模声明
  python q3_verify_grid.py --only overlap --go     # 只做重叠段（快）
  python q3_verify_grid.py --only temporal --go    # 只做时间收敛
  python q3_verify_grid.py --go                    # 两项都做
"""
import sys
import os
import time
import argparse

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass
for _v in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS'):
    os.environ.setdefault(_v, '1')

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), 'Q2', 'code'))

import openpyxl                                                     # noqa: E402
import q2_core_c as base                                            # noqa: E402
import q3_core as q3                                                # noqa: E402
import q3_core_nu as nu                                             # noqa: E402

WS = os.path.abspath(os.path.join(HERE, '..', '..', '..'))
ATT1 = os.path.join(WS, '10_赛题', 'A题', '附件', '附件1.xlsx')
LOG = os.path.join(HERE, 'q3_verify_grid_log.txt')
DT_OUT = 60.0
R0, HC, KM = base.DEF2['R0'], base.DEF2['h'], base.DEF2['km']
BUF = []


def say(s=''):
    print(s)
    BUF.append(s)


def load_env():
    wb = openpyxl.load_workbook(ATT1, data_only=True)
    ws = wb.active
    rows = [r for r in ws.iter_rows(values_only=True)][1:]
    d = [r for r in rows if r[0] is not None]
    t = np.array([float(r[0]) for r in d])
    T = np.array([float(r[1]) for r in d])
    C = np.array([float(r[2]) for r in d])
    wb.close()
    return q3.env_tables_from_att1(t, T, C)


def run_nu26(hh, t_end_s, ts, Tv, Cv, snap_at=()):
    g, _ = nu.geom_refined_surface(R0, dr0=2.5e-4, region=1.0e-3, M=26)
    nsub = int(round(DT_OUT / hh))
    nsteps = int(round(t_end_s / DT_OUT))
    T = np.full(g['NC'], base.DEF2['T0'])
    C = np.full(g['NC'], base.DEF2['C0'])
    snaps = {}
    for n in range(1, nsteps + 1):
        T, C, _ = nu.step_imex_gen(g, hh, nsub, R0, HC, KM, T, C,
                                   ts, Tv, Cv, (n - 1) * DT_OUT)
        if n in snap_at:
            snaps[n] = C.copy()
    return g, C, snaps


def do_overlap(ts, Tv, Cv):
    say()
    say('② Q2 重叠段回归（nu26，0–3 h，逐时刻 C(0) vs result2）')
    _, _, snaps = run_nu26(1.0 / 32.0, 10800.0, ts, Tv, Cv, snap_at=(30, 60, 120, 180))
    wb2 = openpyxl.load_workbook(os.path.join(WS, '20_交付包', '09_代码与复现',
                                              'results', 'result2.xlsx'), read_only=True)
    wsc = wb2['水分浓度']
    r2 = [r for r in wsc.iter_rows(values_only=True)]
    t2v = np.array([float(r[0]) for r in r2[1:]])
    A2 = np.array([[float(x) for x in r[1:]] for r in r2[1:]])
    wb2.close()
    worst = 0.0
    for n in (30, 60, 120, 180):
        tt = n * DT_OUT
        k2 = int(np.argmin(np.abs(t2v - tt)))
        diff = snaps[n][0] - A2[k2, 0]
        worst = max(worst, abs(diff))
        say('   t=%6.0f s  nu26 C(0)=%.6f  Q2 C(0)=%.6f  差 %+.3e'
            % (tt, snaps[n][0], A2[k2, 0], diff))
    say('   最大逐点差 = %.3e（result2 为 4 位小数存储 ⟹ 应 ≤5e-5）' % worst)


def do_temporal(ts, Tv, Cv):
    say()
    say('① 细网格上时间收敛（nu26，24 h，内部步 1/32 vs 1/64）')
    t0 = time.time()
    g, C32, _ = run_nu26(1.0 / 32.0, 86400.0, ts, Tv, Cv)
    say('   1/32 s 完成，耗时 %.1f s' % (time.time() - t0))
    t0 = time.time()
    _, C64, _ = run_nu26(1.0 / 64.0, 86400.0, ts, Tv, Cv)
    say('   1/64 s 完成，耗时 %.1f s' % (time.time() - t0))
    d = np.abs(C64 - C32) / np.maximum(np.abs(C64), 1e-12)
    sel5 = [int(np.argmin(np.abs(g['r'] - k * 1e-3))) for k in (0, 5, 10, 15)] + [g['NC'] - 1]
    for nm, i in zip(('0.0cm', '0.5cm', '1.0cm', '1.5cm', '2.0cm'), sel5):
        say('   %s  C(1/32)=%.6f  C(1/64)=%.6f  相对差 %.3e' % (nm, C32[i], C64[i], d[i]))
    say('   全场最大相对差 = %.3e；判据 5e-5：%s'
        % (float(np.max(d)), 'PASS' if float(np.max(d)) < 5e-5 else '⚠ 未达（细网格需更小步长）'))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--only', choices=['overlap', 'temporal', 'both'], default='both')
    ap.add_argument('--go', action='store_true')
    args = ap.parse_args()
    if not (args.go or os.environ.get('ALLOW_RUN') == '1'):
        print('[SCALE] q3_verify_grid 规模声明：')
        print('        ② overlap ：nu26 至 3 h（≈8 s）')
        print('        ① temporal：nu26 至 24 h，1/32 vs 1/64（≈72 s ＋ 135 s）')
        print('        ⟹ 合计 ≈3.6 min；**不产交付物**。加 --go 执行。')
        return

    ts, Tv, Cv = load_env()
    say('=' * 78)
    say('Q3 · 表面加密网格收尾验证（only=%s）' % args.only)
    say('=' * 78)
    if args.only in ('overlap', 'both'):
        do_overlap(ts, Tv, Cv)
    if args.only in ('temporal', 'both'):
        do_temporal(ts, Tv, Cv)
    with open(LOG, 'a', encoding='utf-8') as f:
        f.write('\n'.join(BUF) + '\n\n')
    print('已追加：%s' % LOG)


if __name__ == '__main__':
    main()
