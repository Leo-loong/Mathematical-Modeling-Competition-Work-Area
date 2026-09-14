# -*- coding: utf-8 -*-
"""
result1.xlsx 一致性核验
=======================
比对「新版单遍结果」与「v3 两遍计算结果（备份）」是否逐值一致。
用于关闭《A_Q1求解记录》§九-2 的异源风险项。

用法：python q1_consistency_check.py
产出：打印比对结论（工作表名 / 维度 / 逐值最大差 / 超容差单元数）
"""
import os
import io
import sys
import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
# 自适应定位工作区根：向上查找含 10_赛题 的目录
# （兼容两种目录深度：工作区 11_建模/11-3_算法与管线/Q1/code/ 与交付包 09 的 code/）
def _find_root(p, _marker='10_赛题', _max=6):
    cur = p
    for _ in range(_max):
        if os.path.isdir(os.path.join(cur, _marker)):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return os.path.abspath(os.path.join(p, '..', '..', '..'))


ROOT = _find_root(HERE)
NEW = os.path.join(ROOT, '20_交付包', '09_代码与复现', 'results', 'result1.xlsx')
OLD = os.path.join(HERE, '_prev_result1.xlsx')
LOG = os.path.join(HERE, 'q1_consistency_log.txt')
TOL = 1e-12

L = []


def say(s):
    L.append(str(s))
    print(s)


def read(fn):
    wb = load_workbook(fn, data_only=True, read_only=True)
    data = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        arr = np.array([[float(v) for v in r[1:]] for r in rows[1:]], dtype=float)
        data[sn] = arr
    return list(wb.sheetnames), data


def main():
    if not os.path.exists(OLD):
        say('缺少备份文件 %s —— 无法比对' % OLD)
        return 1
    sn_new, dnew = read(NEW)
    sn_old, dold = read(OLD)
    say('新版工作表: %s' % sn_new)
    say('旧版工作表: %s' % sn_old)
    ok = (sn_new == sn_old)
    say('工作表名一致 = %s' % ok)
    allok = ok
    for sn in sn_new:
        a, b = dnew[sn], dold[sn]
        if a.shape != b.shape:
            say('[%s] FAIL 维度不同 %s vs %s' % (sn, a.shape, b.shape))
            allok = False
            continue
        d = np.abs(a - b)
        bad = int((d > TOL).sum())
        say('[%s] shape=%s  最大逐值差=%.3e  超容差(>%g)单元数=%d'
            % (sn, a.shape, d.max(), TOL, bad))
        if bad:
            idx = np.argwhere(d > TOL)[:5]
            for i, j in idx:
                say('    差异样例 row=%d col=%d  new=%.10f old=%.10f' % (i + 1, j, a[i, j], b[i, j]))
            allok = False
    say('')
    say('一致性结论：%s' % ('PASS（两版结果逐值一致）' if allok else 'FAIL（存在差异，需查明）'))
    with io.open(LOG, 'w', encoding='utf-8') as f:
        f.write('\n'.join(L) + '\n')
    return 0 if allok else 1


if __name__ == '__main__':
    sys.exit(main())
