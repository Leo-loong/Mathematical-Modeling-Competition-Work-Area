# -*- coding: utf-8 -*-
"""
Q3 · 变步长核 `q3_core_nu` 的**等价性验证 ＋ T1 冒烟 ＋ 几何体检**（只做验证，不产交付物）
================================================================================
目的（P1／P2／F-C 的前置闸门）：
  ① **T1 冒烟**（标 SMOKE）：非均匀核能否跑通（节点少、步数小、限时）；
  ② **等距复现**：同一输入下，`q3_core_nu.step_imex_gen`（在**等距**网格上）
     与已交付核 `q3_core.step_imex_fast` 的逐点最大偏差 **应 ≤1e-11**；
  ③ **回归**：非均匀核在**等距**网格上跑 Q2 重叠段（前 3 h）——与 Q2 的 C(0) 逐位比对；
  ④ **几何体检**：表面加密网格的单元尺寸分布、输出点是否严格落在节点、δ/Δr 判据。

⚠ 本脚本**不产出任何交付数值**；仅打印诊断。

用法：
  python q3_verify_nu.py            # 只看规模声明（默认）
  python q3_verify_nu.py --go       # 真正执行（或设 ALLOW_RUN=1）
"""
import sys
import os
import time
import argparse

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass
for _v in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS'):
    os.environ.setdefault(_v, '1')

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), 'Q2', 'code'))

import openpyxl                                                     # noqa: E402
import q2_core_c as base                                            # noqa: E402
import q3_core as q3                                                # noqa: E402
import q3_core_nu as nu                                             # noqa: E402

WS = os.path.abspath(os.path.join(HERE, '..', '..', '..'))
ATT1 = os.path.join(WS, '10_赛题', 'A题', '附件', '附件1.xlsx')
LOG = os.path.join(HERE, 'q3_verify_nu_log.txt')

R0 = base.DEF2['R0']
HC = base.DEF2['h']
KM = base.DEF2['km']
DT_OUT = 60.0
HH = 1.0 / 32.0
NSUB = int(round(DT_OUT / HH))
BUF = []


def say(s=''):
    print(s)
    BUF.append(s)


def load_env():
    wb = openpyxl.load_workbook(ATT1, data_only=True)
    ws = wb.active
    rows = [r for r in ws.iter_rows(values_only=True)][1:]
    d = [r for r in rows if r[0] is not None]
    t = np.array([float(r[0]) for r in d])
    T = np.array([float(r[1]) for r in d])
    C = np.array([float(r[2]) for r in d])
    wb.close()
    return q3.env_tables_from_att1(t, T, C)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--go', action='store_true')
    args = ap.parse_args()
    if not (args.go or os.environ.get('ALLOW_RUN') == '1'):
        print('[SCALE] q3_verify_nu 规模声明：')
        print('        T1 冒烟  N=40  1 输出步（1920 内部步）  ≤20 s  标 SMOKE')
        print('        等距复现 N=80  20 输出步（38400 内部步）≈6 s')
        print('        回归 3 h N=80  180 输出步（345600 内部步）≈55 s')
        print('        ⟹ 合计 ≈1.5 min；**不产交付物**。加 --go 或 ALLOW_RUN=1 执行。')
        return

    ts, Tv, Cv = load_env()
    say('=' * 78)
    say('Q3 变步长核验证（q3_core_nu）')
    say('=' * 78)

    # ---------- ① T1 冒烟（SMOKE） ----------
    say()
    say('① [SMOKE] 非均匀核跑通测试（N=40，1 输出步）')
    g40 = nu.geom_uniform(40, R0)
    T = np.full(41, base.DEF2['T0'])
    C = np.full(41, base.DEF2['C0'])
    t0 = time.time()
    T, C, info = nu.step_imex_gen(g40, HH, NSUB, R0, HC, KM, T, C, ts, Tv, Cv, 0.0)
    say('   PASS 非均匀核可运行；耗时 %.2f s（含首次 JIT 编译）' % (time.time() - t0))
    say('   C(0)=%.6f  C(R)=%.6f  inner_tot=%d  res_max=%.3e'
        % (C[0], C[-1], info['inner_tot'], info['res_max']))

    # ---------- ② 等距复现 ----------
    say()
    say('② 等距网格上逐点复现已交付核（N=80，20 输出步）')
    g80 = nu.geom_uniform(80, R0)
    T_ref = np.full(81, base.DEF2['T0'])
    C_ref = np.full(81, base.DEF2['C0'])
    T_nu = T_ref.copy()
    C_nu = C_ref.copy()
    dT = dC = 0.0
    t0 = time.time()
    for n in range(1, 21):
        T_ref, C_ref, _ = q3.step_imex_fast(80, 2.5e-4, HH, NSUB, R0, HC, KM,
                                            T_ref, C_ref, ts, Tv, Cv, (n - 1) * DT_OUT)
        T_nu, C_nu, _ = nu.step_imex_gen(g80, HH, NSUB, R0, HC, KM,
                                         T_nu, C_nu, ts, Tv, Cv, (n - 1) * DT_OUT)
        dT = max(dT, float(np.max(np.abs(T_nu - T_ref))))
        dC = max(dC, float(np.max(np.abs(C_nu - C_ref))))
    say('   max|ΔT| = %.3e ℃    max|ΔC| = %.3e kg/kg   （耗时 %.1f s）'
        % (dT, dC, time.time() - t0))
    say('   ⟹ 判定（阈值 1e-11）：%s' % ('PASS' if max(dT, dC) <= 1e-11 else '*** 超阈值，需排查 ***'))

    # ---------- ③ Q2 重叠段回归（3 h） ----------
    say()
    say('③ 非均匀核（等距网格）在 Q2 重叠段（前 3 h）的 C(0) 与 result2 比对')
    T_r = np.full(81, base.DEF2['T0'])
    C_r = np.full(81, base.DEF2['C0'])
    snaps = {}
    for n in range(1, 181):
        T_r, C_r, _ = nu.step_imex_gen(g80, HH, NSUB, R0, HC, KM,
                                       T_r, C_r, ts, Tv, Cv, (n - 1) * DT_OUT)
        if n in (30, 60, 120, 180):
            snaps[n] = C_r[0]
    wb2 = openpyxl.load_workbook(os.path.join(WS, '20_交付包', '09_代码与复现',
                                              'results', 'result2.xlsx'), read_only=True)
    wsc = wb2['水分浓度']
    r2 = [r for r in wsc.iter_rows(values_only=True)]
    t2v = np.array([float(r[0]) for r in r2[1:]])
    A2 = np.array([[float(x) for x in r[1:]] for r in r2[1:]])
    wb2.close()
    worst = 0.0
    for n in (30, 60, 120, 180):
        tt = n * DT_OUT
        k2 = int(np.argmin(np.abs(t2v - tt)))
        diff = snaps[n] - A2[k2, 0]
        worst = max(worst, abs(diff))
        say('   t=%6.0f s   nuC(0)=%.6f   Q2C(0)=%.6f   差 %+.3e'
            % (tt, snaps[n], A2[k2, 0], diff))
    say('   ⟹ 最大逐点差 = %.3e（应 ≤ 输出精度 1e-4）' % worst)

    # ---------- ④ 几何体检 ----------
    say()
    say('④ 表面加密网格几何体检')
    say('   %-6s %-7s %-12s %-12s %-10s %-12s %-12s'
        % ('M', 'N', 's_min/mm', 's_max/mm', 'q', 'δ/s_min@0.15', 'δ/s_min@0.052'))
    for M in (20, 26, 30):
        g, meta = nu.geom_refined_surface(R0, dr0=2.5e-4, region=1.0e-3, M=M)
        d15 = 2.4e-3 * np.exp(-0.45 / 0.15) * np.exp(-3850.0 / 323.15) / KM
        d05 = 2.4e-3 * np.exp(-0.45 / 0.052) * np.exp(-3850.0 / 323.15) / KM
        say('   %-6d %-7d %-12.6f %-12.6f %-10.5f %-12.2f %-12.2f'
            % (M, g['NC'], meta['s_min'] * 1e3, meta['s_max'] * 1e3, meta['q'],
               d15 / meta['s_min'], d05 / meta['s_min']))
    # 输出点是否严格落在节点上
    g, meta = nu.geom_refined_surface(R0, M=26)
    idx = [4 * k for k in range(20)] + [g['NC'] - 1]
    err = max(abs(g['r'][i] - (k * 1e-3 if k < 20 else R0))
              for k, i in enumerate(idx))
    say('   输出点最大定位误差 = %.3e m （应 ≈0）' % err)
    say()
    say('   说明：δ=D(C)/k_m 为表面扩散边界层厚度尺度；δ/s_min ≥ 5 表示该薄层已被分辨。')

    with open(LOG, 'w', encoding='utf-8') as f:
        f.write('\n'.join(BUF) + '\n')
    print('已写出：%s' % LOG)


if __name__ == '__main__':
    main()
