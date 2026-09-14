# -*- coding: utf-8 -*-
"""Q4 W6 lib: data + event-driven solve."""
import os
import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WS = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(HERE))))
N0 = 160
CC = 0.15


def load_data():
    att1 = os.path.join(WS, '10_赛题', 'A题', '附件', '附件1.xlsx')
    wb = openpyxl.load_workbook(att1, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    # ★修复：isinstance 过滤已剔除表头行，原先多余的 [1:] 会误删 t=0 数据行
    #   （附件2 丢 (0,2.0) 后 R(0) 被端点钳位成 1.873 cm，-6.4%）
    rr = [(float(a), float(b), float(c)) for a, b, c in ws.iter_rows(values_only=True)
          if a is not None and not isinstance(a, str)]
    t1 = np.array([x[0] for x in rr])
    T1 = np.array([x[1] for x in rr])
    C1 = np.array([x[2] for x in rr])
    ts_env = np.concatenate([t1, [t1[-1] + 1e-6, t1[-1] + 1.0]])
    Tv = np.concatenate([T1, [50.00, 50.00]])
    Cv = np.concatenate([C1, [0.04999, 0.04999]])
    p2 = os.path.join(WS, '10_赛题', 'A题', '附件', '附件2.xlsx')
    wb2 = openpyxl.load_workbook(p2, read_only=True, data_only=True)
    ws2 = wb2[wb2.sheetnames[0]]
    rr2 = [(float(a), float(b)) for a, b in ws2.iter_rows(values_only=True)
           if a is not None and not isinstance(a, str)]
    ts_R = np.array([x[0] for x in rr2])
    Rv = np.array([x[1] / 100.0 for x in rr2])
    return ts_env, Tv, Cv, ts_R, Rv


def rho_cp(C):
    return (760.0 + 90.0 * C) * (1850.0 + 2150.0 * C / (C + 1.0))

def solve_to_dry(N, n_sub, t_corr=1, d_fac=1.0, hcv=-1.0, kmv=-1.0,
                 dT=0.0, dC=0.0, L=0.0, breaker_h=150.0, hist_every=0,
                 ts_env=None, Tv=None, Cv=None, ts_R=None, Rv=None,
                 verbose=False, tag=''):
    """Event-driven solve to t_dry with substep bisection; optional history."""
    import time
    import q4_core as qc
    dxi = 1.0 / N
    DT = 1.0
    hsub = DT / n_sub
    km = qc.DEF4['km'] if kmv < 0 else kmv
    hc = qc.DEF4['h'] if hcv < 0 else hcv
    Tvv = Tv + dT
    Cvv = Cv * (1.0 + dC)
    T = np.full(N + 1, qc.DEF4['T0'])
    C = np.full(N + 1, qc.DEF4['C0'])
    prevT, prevC, prevt = T.copy(), C.copy(), 0.0
    t_dry = None
    dry = None
    hist = []
    tot_it = 0
    t0w = time.perf_counter()
    max_steps = int(breaker_h * 3600)
    xig = np.arange(N + 1) * dxi
    k = 0
    while k < max_steps:
        k += 1
        t0 = (k - 1) * DT
        T, C, info = qc.step_imex4(N, dxi, hsub, n_sub, qc.DEF4['R0'], hc, km,
                                   T, C, t0, ts_env, Tvv, Cvv, ts_R, Rv,
                                   t_corr=t_corr, d_fac=d_fac, L=L)
        tot_it += info['inner_tot']
        tk = k * DT
        if hist_every and (k % hist_every == 0):
            Rk = float(np.interp(tk, ts_R, Rv))
            W = Rk * Rk * float(np.trapezoid(C * xig, dx=dxi))
            U = Rk * Rk * float(np.trapezoid(rho_cp(C) * T * xig, dx=dxi))
            hist.append((tk, Rk, W, U, float(C[-1]), float(T[-1])))
        if C.max() < CC:
            t_dry = tk
            dry = (T.copy(), C.copy())
            break
        prevT, prevC, prevt = T.copy(), C.copy(), tk
        if verbose and k % 12000 == 0:
            print('  [%s] step %d t=%.1f h C0=%.4f wall=%.0f' %
                  (tag, k, tk / 3600.0, C[0], time.perf_counter() - t0w))
    if t_dry is None:
        return dict(t_dry=None, wall=time.perf_counter() - t0w, it=tot_it, hist=hist)
    t_bad, Tb, Cb = prevt, prevT, prevC
    t_ok, Tok, Cok = float(t_dry), dry[0], dry[1]
    while t_ok - t_bad > 1e-3:
        mid = 0.5 * (t_ok + t_bad)
        w = mid - t_bad
        ns = max(1, int(round(w * n_sub)))
        Tm, Cm, info = qc.step_imex4(N, dxi, w / ns, ns, qc.DEF4['R0'], hc, km,
                                     Tb, Cb, t_bad, ts_env, Tvv, Cvv, ts_R, Rv,
                                     t_corr=t_corr, d_fac=d_fac, L=L)
        tot_it += info['inner_tot']
        if Cm.max() < CC:
            t_ok, Tok, Cok = mid, Tm, Cm
        else:
            t_bad, Tb, Cb = mid, Tm, Cm
    return dict(t_dry=t_ok, wall=time.perf_counter() - t0w, it=tot_it,
                Tend=Tok, Cend=Cok, hist=hist)
