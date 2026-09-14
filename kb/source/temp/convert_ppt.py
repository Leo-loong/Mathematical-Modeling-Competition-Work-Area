# -*- coding: utf-8 -*-
"""convert_ppt.py -- 第二批 PPT/PPTX -> Markdown（输出到「第二批提取后」）。

- .pptx : python-pptx 直接抽取（标题/正文/表格/备注）
- .ppt  : PowerPoint COM 先另存为 .pptx（缓存目录），再走同一抽取逻辑
输出编码 UTF-8；进度写 temp/ppt_progress.txt，报告写 temp/第二批PPT转换报告.txt。
"""
import os
import io
import re
import sys
import time

TMP = os.path.dirname(os.path.abspath(__file__))
SOURCE = os.path.dirname(TMP)
SRC = os.path.join(SOURCE, "第二批")
DST = os.path.join(SOURCE, "第二批提取后")
CACHE = os.path.join(TMP, "b2_pptx_cache")
PROGRESS = os.path.join(TMP, "ppt_progress.txt")
REPORT = os.path.join(TMP, "第二批PPT转换报告.txt")

log = []


def say(m):
    log.append(m)
    print(m)
    try:
        with io.open(PROGRESS, "w", encoding="utf-8") as f:
            f.write("\n".join(log[-60:]))
    except Exception:
        pass


def out_path(src):
    rel = os.path.relpath(src, SRC)
    return os.path.join(DST, os.path.splitext(rel)[0] + ".md")


def write_md(dst, body):
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    with io.open(dst, "w", encoding="utf-8") as f:
        f.write(body)


def shape_text(sh):
    """抽取形状文本；表格按 Markdown 表格输出。"""
    lines = []
    if getattr(sh, "has_table", False) and sh.has_table:
        try:
            tbl = sh.table
            rows = []
            for r in tbl.rows:
                rows.append([c.text.replace("|", "\\|").replace("\n", " ").strip()
                             for c in r.cells])
            if rows:
                head = rows[0]
                lines.append("| " + " | ".join(head) + " |")
                lines.append("|" + "|".join(["---"] * len(head)) + "|")
                for r in rows[1:]:
                    lines.append("| " + " | ".join(r) + " |")
            return lines
        except Exception:
            pass
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            t = "".join(run.text for run in p.runs).strip()
            if not t:
                t = (p.text or "").strip()
            if t:
                lines.append(t)
    return lines


def pptx_to_md(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append("<!-- 第 %d 页 -->" % i)
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except Exception:
            title = ""
        if title:
            out.append("## " + title)
        for sh in slide.shapes:
            for t in shape_text(sh):
                out.append(t)
        try:
            if slide.has_notes_slide:
                nt = (slide.notes_slide.notes_text_frame.text or "").strip()
                if nt:
                    out.append("> 备注: " + nt.replace("\n", " "))
        except Exception:
            pass
        out.append("")
    body = "\n".join(out).strip() + "\n"
    return body


def legacy_one(app, src, i):
    """单个 .ppt：COM 另存为 .pptx；失败返回 None（由调用方降级）。"""
    os.makedirs(CACHE, exist_ok=True)
    safe = re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(src))[:60]
    dst = os.path.join(CACHE, "%04d_%s.pptx" % (i, safe))
    try:
        pres = app.Presentations.Open(src, ReadOnly=True, Untitled=False, WithWindow=False)
        pres.SaveAs(dst, 24)   # 24 = ppSaveAsOpenXMLPresentation
        pres.Close()
        return dst
    except Exception:
        return None


def main():
    t0 = time.time()
    os.makedirs(DST, exist_ok=True)
    pptx_list, ppt_list = [], []
    for dp, _dn, fn in os.walk(SRC):
        for f in fn:
            if f.startswith("~$"):
                continue
            e = os.path.splitext(f)[1].lower()
            p = os.path.join(dp, f)
            if e == ".pptx":
                pptx_list.append(p)
            elif e == ".ppt":
                ppt_list.append(p)
    pptx_list.sort()
    ppt_list.sort()
    say("扫描: pptx %d, ppt %d" % (len(pptx_list), len(ppt_list)))

    ok = fail = 0
    failed = []

    def emit(p, body, why_ok=True, err=""):
        nonlocal ok, fail
        if why_ok and len(body.strip()) >= 20:
            write_md(out_path(p), body)
            ok += 1
        else:
            fail += 1
            failed.append((os.path.relpath(p, SRC), err or "正文过短"))

    # 1) .pptx：python-pptx，并发秒出
    from concurrent.futures import ThreadPoolExecutor, as_completed
    n = len(pptx_list)
    done = 0
    with ThreadPoolExecutor(max_workers=8) as ex:
        futs = {ex.submit(pptx_to_md, p): p for p in pptx_list}
        for fu in as_completed(futs):
            p = futs[fu]
            done += 1
            try:
                emit(p, fu.result())
            except Exception as e:
                emit(p, "", False, repr(e)[:100])
            if done % 20 == 0 or done == n:
                say("pptx 进度 %d/%d" % (done, n))

    # 2) .ppt：PowerPoint COM 逐个转码并立即产出，失败降级到 ppt97 直抽
    if ppt_list:
        import pythoncom
        import win32com.client as win32
        import ppt97
        pythoncom.CoInitialize()
        app = win32.Dispatch("PowerPoint.Application")
        com_ok = com_fail = 0
        try:
            for i, p in enumerate(ppt_list):
                real = legacy_one(app, p, i)
                if real:
                    try:
                        body = pptx_to_md(real)
                        if len(body.strip()) >= 20:
                            write_md(out_path(p), body)
                            ok += 1
                            com_ok += 1
                            continue
                    except Exception:
                        pass
                # 降级：纯 Python 直抽
                text, err = ppt97.extract(p)
                if len(text) >= 20:
                    write_md(out_path(p), text + "\n")
                    ok += 1
                    com_ok += 1
                else:
                    fail += 1
                    com_fail += 1
                    failed.append((os.path.relpath(p, SRC), err or "COM 与直抽均无正文"))
                if (i + 1) % 20 == 0 or i + 1 == len(ppt_list):
                    say("ppt 进度 %d/%d (成功 %d)" % (i + 1, len(ppt_list), com_ok))
        finally:
            try:
                app.Quit()
            except Exception:
                pass
            pythoncom.CoUninitialize()
        say("ppt 通道：COM/直抽成功 %d，失败 %d" % (com_ok, com_fail))

    say("")
    say("PPT 转换完成: 成功 %d, 失败 %d, 耗时 %.1f 分钟"
        % (ok, fail, (time.time() - t0) / 60.0))
    say("--- 失败清单 ---")
    for rel, why in failed:
        say("  %s  << %s" % (rel, why))
    with io.open(REPORT, "w", encoding="utf-8") as f:
        f.write("\n".join(log))
    print("PPT DONE ok=%d fail=%d" % (ok, fail))


if __name__ == "__main__":
    main()
